import math
import json
import datetime
import openpyxl
import base64
import pytz
from django.core.files.base import ContentFile
from decimal import Decimal, ROUND_DOWN, InvalidOperation

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.db.models import Q, Sum, Count, Max, F, FloatField, DecimalField, ExpressionWrapper
from django.db.models.functions import TruncDate, TruncMonth, Coalesce
from django.core.paginator import Paginator
from django.utils.dateparse import parse_date
from django.utils import timezone
from datetime import timedelta
from django.db import transaction
from django.urls import reverse # 🌟 นำเข้า reverse เพิ่มเติม (ใส่ไว้บรรทัดบนสุดของไฟล์ร่วมกับอันอื่น หรือใส่ตรงนี้ก็ได้ครับ)

# Models
from .models import ProductionOrder, ProductionOrderMaterial, BOM, BOMItem, Branch, MfgBranch, Salesperson, ProductionStatus, ProductionTeam, DeliveryStatus, Transporter, QCInspectionLog, QCReceiptDoc, QCPassedMedia, QCFailedMedia
from master_data.models import CompanyInfo
from inventory.models import Product, InventoryDoc, StockMovement
from purchasing.models import PurchaseOrder, PurchaseOrderItem, PurchasePreparation
from .forms import BOMForm, BOMItemFormSet
from .models import BlueprintClaimSplit
from .models import LogisticsClaim
from accounting.models import Expense

# ==========================================
# 🌟 Helper Function: แปลงวันที่ DD/MM/YYYY เป็น YYYY-MM-DD
# ==========================================
def parse_thai_date(date_str):
    if not date_str:
        return None
    try:
        return parse_date(date_str)
    except:
        pass
    try:
        day, month, year = map(int, date_str.split('/'))
        return datetime.date(year, month, day)
    except:
        return None

# ==========================================
# 🧠 ระบบคำนวณคอมมิชชัน
# ==========================================
def process_commission_logic(sale_amount, seller, sale_ref_id):
    now = timezone.now()
    sale_amount = Decimal(str(sale_amount))
    target, _ = CompanySalesTarget.objects.get_or_create(year=now.year, month=now.month)
    target.current_sales += sale_amount
    if target.current_sales >= target.target_amount:
        target.is_unlocked = True
    target.save()

    exec_group = SalesGroup.objects.filter(group_type='EXECUTIVE').first()
    if exec_group:
        exec_pool_amount = sale_amount * (exec_group.commission_rate / Decimal('100'))
        executives = Employee.objects.filter(sales_group=exec_group)
        if executives.exists():
            exact_share = exec_pool_amount / Decimal(str(executives.count()))
            share_per_head = exact_share.quantize(Decimal('0.00'), rounding=ROUND_DOWN)
            remainder = exec_pool_amount - (share_per_head * Decimal(str(executives.count())))
            for ex in executives:
                CommissionLog.objects.create(recipient=ex, source_employee=seller, level=1, amount=share_per_head, sale_ref_id=sale_ref_id)
            if remainder > 0:
                exec_group.fund_balance += remainder
                exec_group.save()
                FundTransaction.objects.create(group=exec_group, transaction_type='IN', amount=remainder, description=f"ปัดเศษเข้าสำรองบริหาร จากบิล {sale_ref_id}")

    if not seller or not seller.sales_group: return
    group = seller.sales_group

    total_group_comm = sale_amount * (group.commission_rate / Decimal('100'))
    fund_to_add = Decimal('0.00')

    if group.group_type == 'TEAM':
        exact_leader_share = total_group_comm * (group.share_leader / Decimal('100'))
        leader_share = exact_leader_share.quantize(Decimal('0.00'), rounding=ROUND_DOWN)
        fund_to_add += (exact_leader_share - leader_share)
        leader = Employee.objects.filter(sales_group=group, group_role='LEADER').first()
        if leader:
            CommissionLog.objects.create(recipient=leader, source_employee=seller, level=1, amount=leader_share, sale_ref_id=sale_ref_id)
        else:
            fund_to_add += leader_share

        exact_l1_share = total_group_comm * (group.share_level1 / Decimal('100'))
        l1_members = Employee.objects.filter(sales_group=group, group_role='LEVEL1')
        if l1_members.exists():
            exact_split = exact_l1_share / Decimal(str(l1_members.count()))
            split = exact_split.quantize(Decimal('0.00'), rounding=ROUND_DOWN)
            remainder = exact_l1_share - (split * Decimal(str(l1_members.count())))
            fund_to_add += remainder
            for m in l1_members:
                CommissionLog.objects.create(recipient=m, source_employee=seller, level=2, amount=split, sale_ref_id=sale_ref_id)
        else:
            fund_to_add += exact_l1_share

        exact_l2_share = total_group_comm * (group.share_level2 / Decimal('100'))
        l2_members = Employee.objects.filter(sales_group=group, group_role='LEVEL2')
        if l2_members.exists():
            exact_split = exact_l2_share / Decimal(str(l2_members.count()))
            split = exact_split.quantize(Decimal('0.00'), rounding=ROUND_DOWN)
            remainder = exact_l2_share - (split * Decimal(str(l2_members.count())))
            fund_to_add += remainder
            for m in l2_members:
                CommissionLog.objects.create(recipient=m, source_employee=seller, level=3, amount=split, sale_ref_id=sale_ref_id)
        else:
            fund_to_add += exact_l2_share

        exact_fixed_fund_share = total_group_comm * (group.share_fund / Decimal('100'))
        fund_to_add += exact_fixed_fund_share

    elif group.group_type == 'INDEPENDENT':
        exact_fund_share = total_group_comm * (group.share_fund / Decimal('100'))
        exact_personal_share = total_group_comm - exact_fund_share
        personal_share = exact_personal_share.quantize(Decimal('0.00'), rounding=ROUND_DOWN)
        fund_to_add = exact_fund_share + (exact_personal_share - personal_share)
        CommissionLog.objects.create(recipient=seller, source_employee=seller, level=1, amount=personal_share, sale_ref_id=sale_ref_id)

    if fund_to_add > 0:
        group.fund_balance += fund_to_add
        group.save()
        FundTransaction.objects.create(group=group, transaction_type='IN', amount=fund_to_add, description=f"รับคอมฯและปัดเศษสตางค์ จากบิล {sale_ref_id}")

def get_next_document_number():
    now = timezone.now()
    thai_year = (now.year + 543) % 100
    prefix = f"DLN-{thai_year:02d}{now.strftime('%m')}"
    last_inv = Invoice.objects.filter(code__startswith=prefix).aggregate(Max('code'))['code__max']
    last_pos = POSOrder.objects.filter(code__startswith=prefix).aggregate(Max('code'))['code__max']
    max_seq = 0
    if last_inv:
        try: max_seq = max(max_seq, int(last_inv.split('-')[-1]))
        except: pass
    if last_pos:
        try: max_seq = max(max_seq, int(last_pos.split('-')[-1]))
        except: pass
    return f"{prefix}-{max_seq + 1:03d}"

def get_target_employees(user):
    current_emp = getattr(user, 'employee', None)
    if user.is_superuser: return Employee.objects.all(), "Admin View"
    elif current_emp:
        dept_name = current_emp.department.name if current_emp.department else ""
        if 'บัญชี' in dept_name or 'Accounting' in dept_name: return Employee.objects.all(), "Accounting View"
        rank = current_emp.business_rank.lower() if current_emp.business_rank else ""
        if rank in ['manager', 'director']: return Employee.objects.all(), "Manager View"
        elif rank == 'supervisor':
            if current_emp.department: return Employee.objects.filter(department=current_emp.department), f"Team {current_emp.department.name}"
            else: return Employee.objects.filter(Q(id=current_emp.id) | Q(introducer=current_emp)), "Direct Team"
        else: return Employee.objects.filter(id=current_emp.id), "Self View"
    return Employee.objects.none(), "-"

def get_sales_queryset(model_class, user, target_employees):
    if user.is_superuser: return model_class.objects.all()
    if hasattr(user, 'employee') and user.employee:
        dept_name = user.employee.department.name if user.employee.department else ""
        rank = user.employee.business_rank.lower() if user.employee.business_rank else ""
        if 'บัญชี' in dept_name or rank in ['manager', 'director']: return model_class.objects.all()

    emp_ids = list(target_employees.values_list('id', flat=True))
    return model_class.objects.filter(employee_id__in=emp_ids)

def is_sales_authorized(user):
    if user.is_superuser: return True
    if hasattr(user, 'employee') and user.employee:
        dept = user.employee.department.name if user.employee.department else ''
        if 'ขาย' in dept or 'Sales' in dept: return True
    return False

# ==========================================
# 🌟 ระบบใบสั่งผลิต (Production Order) 🌟
# ==========================================
@login_required
def production_list(request):
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')
    search_q = request.GET.get('q', '')
    search_salesperson = request.GET.get('salesperson', '')
    search_team = request.GET.get('team', '')
    search_status = request.GET.get('status', '')
    search_branch = request.GET.get('branch', '')

    orders = ProductionOrder.objects.select_related(
        'product', 'quotation_ref', 'salesperson',
        'production_team', 'delivery_status', 'transporter', 'branch'
    ).prefetch_related('completed_departments').all().order_by('-id')

    if start_date:
        orders = orders.filter(start_date__gte=start_date)
    if end_date:
        orders = orders.filter(start_date__lte=end_date)
    if search_status:
        orders = orders.filter(status=search_status)
    if search_branch:
        orders = orders.filter(branch_id=search_branch)
    if search_q:
        orders = orders.filter(
            Q(code__icontains=search_q) | Q(customer_name__icontains=search_q) | Q(quotation_ref__code__icontains=search_q)
        )
    if search_salesperson:
        orders = orders.filter(salesperson_id=search_salesperson)
    if search_team:
        orders = orders.filter(production_team_id=search_team)

    active_qs = orders.filter(is_closed=False)
    closed_qs = orders.filter(is_closed=True)

    active_paginator = Paginator(active_qs, 10)
    active_page_num = request.GET.get('active_page', 1)
    active_orders = active_paginator.get_page(active_page_num)

    closed_paginator = Paginator(closed_qs, 10)
    closed_page_num = request.GET.get('closed_page', 1)
    closed_orders = closed_paginator.get_page(closed_page_num)

    url_params = request.GET.copy()
    if 'active_page' in url_params: del url_params['active_page']
    if 'closed_page' in url_params: del url_params['closed_page']
    filter_string = url_params.urlencode()

    branches = MfgBranch.objects.all().order_by('name')
    prod_statuses = ProductionStatus.objects.all().order_by('sequence', 'id')
    prod_teams = ProductionTeam.objects.all().order_by('name')
    deliv_statuses = DeliveryStatus.objects.all().order_by('name')
    transporters = Transporter.objects.all().order_by('name')
    salespersons = Salesperson.objects.all().order_by('name')
    total_departments_count = prod_statuses.count()
    system_status_choices = ProductionOrder.STATUS_CHOICES

    can_add_master_data = False
    is_sales = False

    if request.user.is_superuser:
        can_add_master_data = True
    else:
        current_emp = getattr(request.user, 'employee', None)
        is_manager = False
        is_planner = False

        user_groups = list(request.user.groups.values_list('name', flat=True))
        if any('Manager' in g for g in user_groups): is_manager = True
        if any('Planner' in g or 'Production' in g for g in user_groups): is_planner = True

        if current_emp:
            rank = current_emp.business_rank.lower() if current_emp.business_rank else ""
            job_title = current_emp.position.title.lower() if current_emp.position else ""
            dept_name = current_emp.department.name if current_emp.department else ""

            if rank in ['manager', 'director'] or 'manager' in job_title:
                is_manager = True
            if 'วางแผน' in dept_name or 'ผลิต' in dept_name:
                 is_planner = True

            if 'ขาย' in dept_name or 'sales' in dept_name.lower():
                 is_sales = True

        if is_manager and is_planner:
            can_add_master_data = True

    return render(request, 'manufacturing/production_list.html', {
        'active_orders': active_orders,
        'closed_orders': closed_orders,
        'branches': branches,
        'prod_statuses': prod_statuses,
        'prod_teams': prod_teams,
        'deliv_statuses': deliv_statuses,
        'transporters': transporters,
        'salespersons': salespersons,
        'start_date': start_date,
        'end_date': end_date,
        'search_q': search_q,
        'search_branch': search_branch,
        'search_salesperson': search_salesperson,
        'search_team': search_team,
        'search_status': search_status,
        'system_status_choices': system_status_choices,
        'filter_string': filter_string,
        'can_add_master_data': can_add_master_data,
        'is_sales': is_sales,
        'total_departments_count': total_departments_count,
    })

@login_required
def planner_board(request):
    today = timezone.now().date()
    default_start = today - datetime.timedelta(days=30)
    default_end = today + datetime.timedelta(days=30)

    start_date = request.GET.get('start_date', default_start.strftime('%Y-%m-%d'))
    end_date = request.GET.get('end_date', default_end.strftime('%Y-%m-%d'))

    search_q = request.GET.get('q', '')
    search_branch = request.GET.get('branch', '')
    search_team = request.GET.get('team', '')
    search_salesperson = request.GET.get('salesperson', '')

    # 🌟 [FIXED] เลิกซ่อนการ์ดจากการกด is_closed เปลี่ยนมาซ่อนเมื่อส่งมอบให้ลูกค้าสำเร็จเท่านั้น
    orders = ProductionOrder.objects.select_related(
        'product', 'branch', 'production_team', 'salesperson', 'salesperson__branch', 'quotation_ref', 'delivery_status'
    ).exclude(
        delivery_status__name__in=['ส่งมอบสำเร็จ', 'ลูกค้าเซ็นรับแล้ว', 'จัดส่งเรียบร้อย']
    ).order_by('-id')

    if start_date: orders = orders.filter(start_date__gte=start_date)
    if end_date: orders = orders.filter(start_date__lte=end_date)
    if search_q:
        orders = orders.filter(Q(code__icontains=search_q) | Q(customer_name__icontains=search_q) | Q(product__name__icontains=search_q))
    if search_branch: orders = orders.filter(branch_id=search_branch)
    if search_team: orders = orders.filter(production_team_id=search_team)
    if search_salesperson: orders = orders.filter(salesperson_id=search_salesperson)

    branches = MfgBranch.objects.all().order_by('name')
    teams = ProductionTeam.objects.all().order_by('name')
    salespersons = Salesperson.objects.select_related('branch').all().order_by('name')

    prod_statuses = ProductionStatus.objects.all().order_by('sequence', 'id')
    deliv_statuses = DeliveryStatus.objects.all().order_by('name')
    transporters = Transporter.objects.all().order_by('name')

    return render(request, 'manufacturing/planner_board.html', {
        'orders': orders,
        'branches': branches,
        'teams': teams,
        'salespersons': salespersons,
        'prod_statuses': prod_statuses,
        'deliv_statuses': deliv_statuses,
        'transporters': transporters,
        'start_date': start_date,
        'end_date': end_date,
        'search_q': search_q,
        'search_branch': search_branch,
        'search_team': search_team,
        'search_salesperson': search_salesperson,
    })

@login_required
def inventory_board(request):
    # 🌟 1. ส่วนของแท็บที่ 1: ดึงรายการที่ "รอเบิก"
    orders = ProductionOrder.objects.select_related('quotation_ref', 'salesperson', 'branch', 'production_team', 'product').filter(status='WAITING_INVENTORY', is_closed=False).order_by('status', '-id')
    virtual_stock = {}

    for order in orders:
        shortage = []
        for item in order.materials.all():
            mat_id = item.raw_material.id
            required_qty = Decimal(str(item.quantity))

            if mat_id not in virtual_stock:
                virtual_stock[mat_id] = max(Decimal('0'), item.raw_material.stock_qty or Decimal('0'))

            if virtual_stock[mat_id] < required_qty:
                missing = required_qty - virtual_stock[mat_id]
                shortage.append(f"{item.raw_material.name} (ขาดอีก {missing:.2f})")
                virtual_stock[mat_id] = Decimal('0')
            else:
                virtual_stock[mat_id] -= required_qty

        order.has_shortage = len(shortage) > 0
        order.shortage_list = ", ".join(shortage)

    # 🌟 2. ส่วนของแท็บที่ 2: ดึง "ประวัติการเบิก" ย้อนหลัง
    history_docs = InventoryDoc.objects.select_related('created_by').filter(doc_type='GI', reference__startswith='เบิกผลิต').order_by('-created_at')[:30]

    # 🌟 [NEW] ระบบแกะรหัส JOB เพื่อดึงเลขที่ใบเสนอราคา (QT) มาโชว์ในประวัติ
    job_codes = []
    for doc in history_docs:
        if doc.reference and doc.reference.startswith("เบิกผลิต "):
            job_code = doc.reference.replace("เบิกผลิต ", "").strip()
            job_codes.append(job_code)
            doc.extracted_job_code = job_code
        else:
            doc.extracted_job_code = None

    # นำรหัส JOB ที่แกะได้ ไปค้นหา ProductionOrder เพื่อดึง quotation_ref__code
    jobs_with_qt = ProductionOrder.objects.filter(code__in=job_codes).select_related('quotation_ref')
    job_to_qt_map = {job.code: (job.quotation_ref.code if job.quotation_ref else "ไม่มีอ้างอิง") for job in jobs_with_qt}

    # แปะเลขใบเสนอราคากลับเข้าไปใน history_docs แต่ละบรรทัด
    for doc in history_docs:
        if doc.extracted_job_code:
            doc.qt_code = job_to_qt_map.get(doc.extracted_job_code, "ไม่พบข้อมูล")
        else:
            doc.qt_code = "-"

    return render(request, 'manufacturing/inventory_board.html', {
        'orders': orders,
        'history_docs': history_docs
    })

@login_required
def production_create(request):
    boms = BOM.objects.select_related('product').all()
    fg_with_bom = [bom.product for bom in boms]
    branches = Branch.objects.all().order_by('name')
    salespersons = Salesperson.objects.select_related('branch').all().order_by('name')

    if request.method == 'POST':
        product_id = request.POST.get('product')
        note = request.POST.get('note', '')
        start_date = request.POST.get('start_date')
        delivery_date = request.POST.get('delivery_date')
        branch_id = request.POST.get('branch')
        salesperson_id = request.POST.get('salesperson')
        customer_name = request.POST.get('customer_name', '')

        if product_id:
            product = get_object_or_404(Product, id=product_id)
            is_onsite = request.POST.get('is_onsite') == 'on'
            order = ProductionOrder(
                product=product,
                quantity=1,
                status='NEW_JOB',
                note=note,
                customer_name=customer_name,
                is_onsite=is_onsite,
                responsible_person=getattr(request.user, 'employee', None)
            )

            if start_date: order.start_date = start_date
            if delivery_date: order.delivery_date = delivery_date
            if salesperson_id: order.salesperson_id = salesperson_id

            order.save()
            messages.success(request, f"✅ สร้างใบสั่งผลิต {order.code} สำเร็จ! งานถูกส่งไปยังคอลัมน์รอจ่ายงาน (W1)")
            return redirect('planner_board')
        else:
            messages.error(request, "❌ กรุณาเลือกแบบบ้านที่ต้องการผลิต")

    return render(request, 'manufacturing/production_form.html', {
        'fg_with_bom': fg_with_bom,
        'branches': branches,
        'salespersons': salespersons
    })

@login_required
def start_production(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    if not order.materials.exists():
        messages.warning(request, "⚠️ กรุณาดึงรายการวัตถุดิบ (BOM) ก่อนกดยืนยันแผนการผลิต!")
        return redirect('production_detail', pk=order.pk)

    order.status = 'WAITING_MATERIALS'
    order.save()
    messages.success(request, f"🚀 ยืนยันแผนงาน {order.code} แล้ว! งานถูกส่งไปที่ฝ่ายจัดซื้อ")
    return redirect('production_list')

@login_required
def ppo_prepare(request):
    # กรองไม่ให้ดึง JOB ที่มี ppo_ref แล้วมาแสดงซ้ำ
    available_jobs = ProductionOrder.objects.filter(status='WAITING_MATERIALS', is_materials_ordered=False, ppo_ref__isnull=True).order_by('code')
    ppo_code = ""
    materials_by_supplier = {}
    selected_job_ids = []

    if request.method == 'POST':
        selected_job_ids = request.POST.getlist('job_ids')
        if selected_job_ids:
            ppo = PurchasePreparation.objects.create(created_by=getattr(request.user, 'employee', None))
            jobs = ProductionOrder.objects.filter(id__in=selected_job_ids)
            ppo.production_orders.set(jobs)
            ppo_code = ppo.code

            virtual_stock_used = {}

            for job in jobs:
                # 🌟 [FIXED] เปลี่ยนให้ระบบดึงวัตถุดิบจากใบ JOB จริงๆ แทนการเช็ค BOM ดั้งเดิม 🌟
                for item in job.materials.all():
                    total_needed = Decimal(str(item.quantity))
                    mat_id = item.raw_material.id

                    # คำนวณสต๊อกคงเหลือ
                    actual_stock = max(Decimal('0'), item.raw_material.stock_qty or Decimal('0'))
                    used_stock = virtual_stock_used.get(mat_id, Decimal('0'))
                    available_stock = max(Decimal('0'), actual_stock - used_stock)

                    # ลอจิกหักลบสต๊อก
                    if available_stock >= total_needed:
                        to_order_qty = Decimal('0')
                        virtual_stock_used[mat_id] = used_stock + total_needed
                    else:
                        to_order_qty = total_needed - available_stock
                        virtual_stock_used[mat_id] = used_stock + available_stock

                    supplier = item.raw_material.supplier
                    sup_id = supplier.id if supplier else "none"
                    sup_name = supplier.name if supplier else "ไม่ได้ระบุร้านค้า"

                    if sup_id not in materials_by_supplier:
                        materials_by_supplier[sup_id] = {'name': sup_name, 'items': {}}

                    if mat_id not in materials_by_supplier[sup_id]['items']:
                        materials_by_supplier[sup_id]['items'][mat_id] = {
                            'product_id': item.raw_material.id,
                            'product_name': item.raw_material.name,
                            'product_code': item.raw_material.code,
                            'total_needed': Decimal('0'),
                            'in_stock': actual_stock,
                            'qty': Decimal('0'),
                            'cost': item.raw_material.cost_price or Decimal('0'),
                            'total': Decimal('0')
                        }

                    materials_by_supplier[sup_id]['items'][mat_id]['total_needed'] += total_needed
                    materials_by_supplier[sup_id]['items'][mat_id]['qty'] += to_order_qty
                    materials_by_supplier[sup_id]['items'][mat_id]['total'] = materials_by_supplier[sup_id]['items'][mat_id]['qty'] * materials_by_supplier[sup_id]['items'][mat_id]['cost']

            # อัปเดตสถานะให้ JOB ทุกใบที่เลือก
            jobs.update(is_materials_ordered=True, ppo_ref=ppo_code)

            # คลีนข้อมูล และกรองเอาเฉพาะอันที่ต้องสั่งซื้อ (>0) ส่งไปทำ PO
            suppliers_to_remove = []
            for sup_id in materials_by_supplier:
                items_list = []
                for v in materials_by_supplier[sup_id]['items'].values():
                    if float(v['qty']) > 0:
                        v['total_needed'] = float(v['total_needed'])
                        v['in_stock'] = float(v['in_stock'])
                        v['qty'] = float(v['qty'])
                        v['cost'] = float(v['cost'])
                        v['total'] = float(v['total'])
                        items_list.append(v)

                if items_list:
                    materials_by_supplier[sup_id]['items'] = items_list
                else:
                    suppliers_to_remove.append(sup_id)

            for sup_id in suppliers_to_remove:
                del materials_by_supplier[sup_id]

            # 🌟 [NEW] เพิ่มส่วนนี้: เก็บข้อมูลไว้ใน Session ชั่วคราว เพื่อส่งไปแสดงหน้าถัดไป 🌟
            request.session['ppo_code'] = ppo_code
            request.session['materials_by_supplier'] = materials_by_supplier

            messages.success(request, f"สร้างใบเตรียมการสั่งซื้อสำเร็จ! เลขที่เอกสาร PPO มัดรวม: {ppo_code}")

            # 🌟 [CRITICAL FIX] บังคับ Redirect กลับหน้าเดิม (หรือหน้ารายละเอียด PPO) เพื่อล้างสถานะ POST 🌟
            return redirect('ppo_prepare') # จะรีไดเรกต์กลับมาหน้าตัวเองแบบโหลดใหม่หมดจด

        else:
            messages.warning(request, "⚠️ กรุณาติ๊กเลือกอย่างน้อย 1 ใบสั่งผลิต (JOB) เพื่อคำนวณวัตถุดิบ")

    # 🌟 [NEW] ดึงข้อมูลที่เพิ่งทำเสร็จออกมาโชว์ให้ชื่นใจ 🌟
    display_ppo_code = request.session.pop('ppo_code', '')
    display_materials = request.session.pop('materials_by_supplier', {})

    return render(request, 'manufacturing/ppo_prepare.html', {
        'available_jobs': available_jobs,
        'ppo_code': display_ppo_code,
        'materials_by_supplier': display_materials,
    })

@login_required
@transaction.atomic
def materials_ready(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    job_materials = order.materials.all()
    if not job_materials:
        messages.error(request, "❌ ไม่พบรายการวัตถุดิบในใบสั่งผลิตนี้")
        return redirect('inventory_board')

    shortage = []
    for item in job_materials:
        required_qty = Decimal(str(item.quantity))
        actual_stock = max(Decimal('0'), item.raw_material.stock_qty or Decimal('0'))
        if actual_stock < required_qty:
            shortage.append(f"{item.raw_material.name} (ขาด {required_qty - actual_stock:.2f})")

    if shortage:
        err_msg = " / ".join(shortage)
        messages.error(request, f"❌ ไม่สามารถตัดสต็อกได้! วัตถุดิบในคลังไม่พอ: {err_msg}")
        return redirect('inventory_board')

    # 🌟 [UPDATE] สร้างใบเบิกพร้อมติด Tag 'RM' เพื่อให้รันเลขเอกสารแบบใหม่
    doc_out = InventoryDoc(
        doc_type='GI',
        reference=f"เบิกผลิต {order.code}",
        description=f"เบิกวัตถุดิบเตรียมผลิต {order.product.name}",
        created_by=request.user
    )
    doc_out._doc_category = 'RM' # ส่งซิกว่าเป็น RM (วัตถุดิบ)
    doc_out.save()
    for item in job_materials:
        StockMovement.objects.create(doc=doc_out, product=item.raw_material, quantity=Decimal(str(item.quantity)), movement_type='OUT', created_by=request.user)

    # 🌟 [FIXED] ยกเลิกทางลัด On-site งานทุกประเภทเมื่อเบิกของเสร็จ ต้องไปที่ "กระดานผลิต (IN_PROGRESS)" ทั้งหมด
    order.status = 'IN_PROGRESS'
    order.save()
    messages.success(request, f"✅ ตัดสต็อกสำเร็จ! สถานะเปลี่ยนเป็น 'งานอยู่ระหว่างผลิต' เพื่อรอช่างและ QC ตรวจสอบแล้ว")

    return redirect('inventory_board')

@login_required
def start_actual_production(request, pk):
    return redirect('production_list')

@login_required
@transaction.atomic
def production_process(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    if order.status == 'COMPLETED':
        messages.warning(request, "⚠️ รายการนี้ผลิตเสร็จและรับเข้าคลังไปแล้ว!")
        return redirect('production_list')

    main_product = order.product

    desc_text = f"รับสินค้า {main_product.name} จากการผลิต {order.code}"
    if order.customer_name:
        desc_text += f" (ลค. {order.customer_name})"

    # 🌟 [FIXED] แยกตะกร้ารับเข้าเป็นของฝ่ายผลิตโดยเฉพาะ (GR-FG)
    qc_doc = QCReceiptDoc.objects.create(production_order=order, product=main_product, quantity=Decimal(str(order.quantity)), created_by=request.user)
    StockMovement.objects.create(doc=None, reference_doc=qc_doc.doc_no, note=desc_text, product=main_product, quantity=Decimal(str(order.quantity)), movement_type='IN', created_by=request.user)

    order.status = 'COMPLETED'
    order.finish_date = timezone.now().date()
    order.save()

    messages.success(request, f"🎉 สำเร็จ! รับ {main_product.name} ({order.quantity} หลัง) เข้าคลังเรียบร้อยแล้ว!")
    return redirect('production_list')

# ==========================================
# 🌟 [REVERTED] ฟังก์ชัน production_detail กลับสู่โหมดดูข้อมูลปกติ 🌟
# ==========================================
@login_required
def production_detail(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    materials = order.materials.all()
    has_bom = BOM.objects.filter(product=order.product).exists()
    raw_materials = Product.objects.filter(product_type='RM', is_active=True).order_by('code')

    return render(request, 'manufacturing/production_detail.html', {
        'order': order,
        'materials': materials,
        'has_bom': has_bom,
        'raw_materials': raw_materials
    })

@login_required
def print_bom(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    materials = order.materials.all()
    company = CompanyInfo.objects.first()
    return render(request, 'manufacturing/print_bom.html', {'order': order, 'materials': materials, 'company': company})

@login_required
def upload_blueprint(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    if request.method == 'POST' and 'blueprint_file' in request.FILES:
        order.blueprint_file = request.FILES['blueprint_file']
        order.save()
        messages.success(request, f"📎 แนบไฟล์แบบแปลน (Blueprint) สำหรับ {order.code} เรียบร้อยแล้ว")
    return redirect('production_detail', pk=order.pk)

@login_required
def load_standard_bom(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    referer = request.META.get('HTTP_REFERER')

    if order.materials.exists():
        messages.warning(request, "⚠️ มีการดึงรายการสูตรผลิตในใบสั่งผลิตนี้ไปแล้ว")
        return redirect(referer) if referer else redirect('production_detail', pk=order.pk)

    bom = BOM.objects.filter(product=order.product).first()
    if not bom:
        messages.error(request, "❌ ไม่พบสูตรผลิตมาตรฐาน (BOM) สำหรับสินค้ารุ่นนี้")
        return redirect(referer) if referer else redirect('production_detail', pk=order.pk)

    for item in bom.items.all():
        ProductionOrderMaterial.objects.create(
            production_order=order,
            raw_material=item.raw_material,
            quantity=item.quantity * order.quantity,
            is_additional=False
        )
    messages.success(request, "📋 ดึงรายการวัตถุดิบจากสูตรมาตรฐานเรียบร้อยแล้ว!")
    return redirect(referer) if referer else redirect('production_detail', pk=order.pk)

@login_required
def add_additional_material(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    referer = request.META.get('HTTP_REFERER')

    if request.method == 'POST':
        product_id = request.POST.get('raw_material')
        try:
            qty = Decimal(request.POST.get('quantity', '0'))
        except:
            qty = Decimal('0')

        if product_id and qty > 0:
            rm = get_object_or_404(Product, pk=product_id)
            ProductionOrderMaterial.objects.create(
                production_order=order,
                raw_material=rm,
                quantity=qty,
                is_additional=True
            )
            messages.success(request, f"➕ เพิ่มวัตถุดิบส่วนเพิ่ม: {rm.name} จำนวน {qty} ลงในบิลเรียบร้อยแล้ว")
        else:
            messages.error(request, "❌ กรุณาเลือกวัตถุดิบและใส่จำนวนให้ถูกต้อง")
    return redirect(referer) if referer else redirect('production_detail', pk=order.pk)

@login_required
def delete_production_material(request, pk):
    mat = get_object_or_404(ProductionOrderMaterial, pk=pk)
    order_id = mat.production_order.id
    referer = request.META.get('HTTP_REFERER')

    if mat.production_order.status != 'COMPLETED':
        mat.delete()
        messages.success(request, "🗑️ ลบรายการวัตถุดิบออกจากบิลผลิตเรียบร้อยแล้ว")
    else:
        messages.error(request, "❌ ไม่สามารถลบได้เนื่องจากใบสั่งผลิตนี้ดำเนินการเสร็จสิ้นแล้ว")
    return redirect(referer) if referer else redirect('production_detail', pk=order_id)

@login_required
def blueprint_viewer(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    return render(request, 'manufacturing/blueprint_viewer.html', {'order': order})

@login_required
def generate_pos_from_production(request, pk):
    messages.info(request, "ฟังก์ชันนี้กำลังอยู่ระหว่างการพัฒนาเพิ่มเติม")
    return redirect('ppo_prepare')

@login_required
def production_print(request, po_id):
    po = get_object_or_404(ProductionOrder, id=po_id)
    bom = BOM.objects.filter(product=po.product).first()
    company = CompanyInfo.objects.first()
    return render(request, 'manufacturing/production_print.html', {'po': po, 'bom': bom, 'company': company})

@login_required
def bom_create(request):
    if request.method == 'POST':
        form = BOMForm(request.POST)
        formset = BOMItemFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
             bom = form.save()
             formset.instance = bom
             formset.save()
             messages.success(request, f"✅ สร้างสูตรผลิตสำหรับ {bom.product.name} เรียบร้อยแล้ว!")
             return redirect('bom_detail', pk=bom.pk)
    else:
        form = BOMForm()
        formset = BOMItemFormSet()
    return render(request, 'manufacturing/bom_form.html', {'form': form, 'formset': formset, 'title': 'สร้างสูตรผลิตใหม่ (New BOM)'})

@login_required
def bom_detail(request, pk):
    bom = get_object_or_404(BOM, pk=pk)
    items = bom.items.select_related('raw_material', 'raw_material__supplier').all()

    total_cost = Decimal('0')
    for item in items:
        item.calculated_total_cost = Decimal(str(item.quantity)) * (item.raw_material.cost_price or Decimal('0'))
        total_cost += item.calculated_total_cost

    return render(request, 'manufacturing/bom_detail.html', {'bom': bom, 'items': items, 'total_cost': total_cost})

@login_required
def bom_edit(request, pk):
    bom = get_object_or_404(BOM.objects.prefetch_related('items'), pk=pk)

    if request.method == 'POST':
        form = BOMForm(request.POST, instance=bom)
        formset = BOMItemFormSet(request.POST, instance=bom)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, f"✅ แก้ไขสูตรผลิต {bom.product.name} เรียบร้อยแล้ว!")
            return redirect('bom_detail', pk=bom.pk)
    else:
        form = BOMForm(instance=bom)
        formset = BOMItemFormSet(instance=bom)

    return render(request, 'manufacturing/bom_form.html', {'form': form, 'formset': formset, 'title': f'แก้ไขสูตรผลิต: {bom.product.code}'})

@login_required
def print_master_bom(request, pk):
    bom = get_object_or_404(BOM, pk=pk)
    items = bom.items.select_related('raw_material', 'raw_material__supplier').all()

    total_cost = Decimal('0')
    for item in items:
        item.calculated_total_cost = Decimal(str(item.quantity)) * (item.raw_material.cost_price or Decimal('0'))
        total_cost += item.calculated_total_cost

    company = CompanyInfo.objects.first()
    return render(request, 'manufacturing/print_master_bom.html', {
        'bom': bom,
        'items': items,
        'total_cost': total_cost,
        'company': company
    })

@login_required
def update_production_board(request, pk):
    if request.method == 'POST':
        order = get_object_or_404(ProductionOrder, pk=pk)
        action = request.POST.get('action')
        redirect_to = request.POST.get('redirect_to')

        if action == 'update_note':
            order.note = request.POST.get('note', '')
            order.save()
            messages.success(request, f"💾 บันทึกหมายเหตุสำหรับ {order.code} เรียบร้อยแล้ว")
            if redirect_to == 'inventory_board':
                return redirect('inventory_board')

        if 'completed_departments' in request.POST:
            selected_depts = request.POST.getlist('completed_departments')
            order.completed_departments.set(selected_depts)
        elif action in ['update_progress', 'qc_complete_and_receive']:
             order.completed_departments.clear()

        if 'production_team' in request.POST:
            order.production_team_id = request.POST.get('production_team') or None

        if 'branch' in request.POST:
            order.branch_id = request.POST.get('branch') or None

        if action == 'dispatch' and order.status == 'NEW_JOB':
            order.status = 'WAITING_BLUEPRINT'
            messages.success(request, f"✅ จ่ายงาน {order.code} ไปยัง {order.branch.name if order.branch else 'ไม่ระบุโรงงาน'} เรียบร้อยแล้ว! (สถานะ: ตรวจแบบแปลน)")

        elif action in ['update_progress', 'qc_complete_and_receive']:
            order.qc_paint = request.POST.get('qc_paint') == 'on'
            order.qc_internal = request.POST.get('qc_internal') == 'on'
            order.qc_external = request.POST.get('qc_external') == 'on'
            order.qc_electrical = request.POST.get('qc_electrical') == 'on'
            order.qc_plumbing = request.POST.get('qc_plumbing') == 'on'
            order.qc_aircon = request.POST.get('qc_aircon') == 'on'

            if action == 'qc_complete_and_receive':
                if order.status != 'COMPLETED':
                    main_product = order.product

                    desc_text = f"รับสินค้า {main_product.name} จาก {order.code} (QC Pass)"
                    if order.customer_name:
                        desc_text += f" (ลค. {order.customer_name})"

                    # 🌟 [FIXED] แยกตะกร้ารับเข้าเป็นของฝ่ายผลิตโดยเฉพาะ (GR-FG)
                    qc_doc = QCReceiptDoc.objects.create(production_order=order, product=main_product, quantity=Decimal(str(order.quantity)), created_by=request.user)
                    StockMovement.objects.create(doc=None, reference_doc=qc_doc.doc_no, note=desc_text, product=main_product, quantity=Decimal(str(order.quantity)), movement_type='IN', created_by=request.user)

                    order.status = 'COMPLETED'
                    order.finish_date = timezone.now().date()
                    order.is_qc_passed = True
                    messages.success(request, f"🎉 ผ่าน QC และรับ {main_product.name} เข้าคลังแล้ว! งานย้ายไปช่อง W4 พร้อมส่ง")
            else:
                messages.success(request, f"✅ อัปเดตความคืบหน้า {order.code} เรียบร้อยแล้ว!")

        is_closed = request.POST.get('is_closed')
        if is_closed == 'on':
            order.is_closed = True
            messages.success(request, f"✅ ปิดจ๊อบ {order.code} อย่างสมบูรณ์แล้ว! (ย้ายไปแท็บประวัติ)")
        elif action in ['update_progress', 'qc_complete_and_receive'] and not is_closed:
            if order.status != 'COMPLETED':
                order.is_closed = False

        order.save()

    return redirect(request.META.get('HTTP_REFERER', 'production_list'))

@login_required
def ajax_add_branch(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        name = data.get('name')
        if name:
            obj, created = Branch.objects.get_or_create(name=name)
            return JsonResponse({'success': True, 'id': obj.id, 'name': obj.name})
    return JsonResponse({'success': False})

@login_required
def ajax_add_salesperson(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        name = data.get('name')
        branch_id = data.get('branch_id')
        if name and branch_id:
            obj, created = Salesperson.objects.get_or_create(name=name, branch_id=branch_id)
            return JsonResponse({'success': True, 'id': obj.id, 'name': obj.name, 'branch_id': obj.branch.id})
    return JsonResponse({'success': False})

@login_required
def ajax_add_prod_status(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        name = data.get('name')
        if name:
            obj, _ = ProductionStatus.objects.get_or_create(name=name, defaults={'sequence': 99})
            return JsonResponse({'success': True, 'id': obj.id, 'name': obj.name})
    return JsonResponse({'success': False})

@login_required
def ajax_add_prod_team(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        name = data.get('name')
        if name:
            obj, _ = ProductionTeam.objects.get_or_create(name=name)
            return JsonResponse({'success': True, 'id': obj.id, 'name': obj.name})
    return JsonResponse({'success': False})

@login_required
def ajax_add_delivery_status(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        name = data.get('name')
        if name:
            obj, _ = DeliveryStatus.objects.get_or_create(name=name)
            return JsonResponse({'success': True, 'id': obj.id, 'name': obj.name})
    return JsonResponse({'success': False})

@login_required
def ajax_add_transporter(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        name = data.get('name')
        if name:
            obj, _ = Transporter.objects.get_or_create(name=name)
            return JsonResponse({'success': True, 'id': obj.id, 'name': obj.name})
    return JsonResponse({'success': False})

@login_required
def ajax_get_fg_by_category(request):
    category_id = request.GET.get('category_id')
    products = Product.objects.filter(product_type='FG', is_active=True).exclude(code__contains='-JOB')
    if category_id:
        products = products.filter(category_id=category_id)
    product_list = list(products.values('id', 'name', 'code'))
    return JsonResponse({'products': product_list})

# ==========================================
# 🌟 [NEW] ดาวน์โหลดไฟล์แม่แบบ Excel สำหรับ BOM 🌟
# ==========================================
@login_required
def download_bom_template(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    # สร้างไฟล์ Excel เปล่า
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM_Template"

    # กำหนดหัวคอลัมน์
    headers = ['FG SKU', 'RM SKU', 'Quantity']
    ws.append(headers)

    # ตกแต่งหัวคอลัมน์ให้สวยงาม (สีพื้นหลังเขียว, ตัวหนา, จัดกึ่งกลาง)
    header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    align_center = Alignment(horizontal="center", vertical="center")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # ปรับความกว้างคอลัมน์
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15

    # ใส่ข้อมูลตัวอย่าง 1 บรรทัด (เพื่อให้ผู้ใช้เข้าใจรูปแบบ)
    ws.append(['ตัวอย่าง: SD169', 'ตัวอย่าง: RM-6910-001', 1])

    # ทำให้แถวข้อมูลตัวอย่างเป็นตัวเอียงสีเทา
    sample_font = Font(italic=True, color="808080")
    for col_num in range(1, 4):
        ws.cell(row=2, column=col_num).font = sample_font

    # เตรียมส่งไฟล์ให้ผู้ใช้ดาวน์โหลด
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="BOM_Import_Template.xlsx"'
    wb.save(response)

    return response

# (ฟังก์ชัน import_bom_excel เดิมของคุณปอนด์จะอยู่ตรงนี้ครับ)

@login_required
def import_bom_excel(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        try:
            df = pd.read_excel(excel_file)
            required_columns = ['FG SKU', 'RM SKU', 'Quantity']
            if not all(col in df.columns for col in required_columns):
                messages.error(request, "❌ รูปแบบไฟล์ไม่ถูกต้อง กรุณาใช้คอลัมน์: FG SKU, RM SKU, Quantity")
                return redirect('import_bom_excel')
            with transaction.atomic():
                 import_count = 0
                 for _, row in df.iterrows():
                    fg_sku = str(row['FG SKU']).strip()
                    rm_sku = str(row['RM SKU']).strip()
                    try:
                        qty = Decimal(str(row['Quantity']))
                    except (ValueError, TypeError, InvalidOperation):
                        continue
                    fg_product = Product.objects.filter(code=fg_sku, product_type='FG').first()
                    if not fg_product:
                        continue
                    bom_obj, _ = BOM.objects.get_or_create(
                        product=fg_product,
                        defaults={'name': f"สูตรมาตรฐาน - {fg_product.name}"}
                     )
                    rm_product = Product.objects.filter(code=rm_sku, product_type='RM').first()
                    if rm_product:
                        BOMItem.objects.update_or_create(
                            bom=bom_obj,
                            raw_material=rm_product,
                            defaults={'quantity': qty}
                        )
                        import_count += 1
            messages.success(request, f"✅ นำเข้าข้อมูลสูตรผลิตสำเร็จทั้งหมด {import_count} รายการ")
            return redirect('import_bom_excel')
        except Exception as e:
            messages.error(request, f"❌ เกิดข้อผิดพลาด: {str(e)}")
    return render(request, 'manufacturing/import_bom.html')

@login_required
def ajax_search_raw_material(request):
    q = request.GET.get('q', '')
    products = Product.objects.filter(product_type='RM', is_active=True)
    if q: products = products.filter(Q(code__icontains=q) | Q(name__icontains=q))
    products = products.order_by('code')[:30]

    # 🌟 [UPDATE] แนบค่า cost_price ส่งกลับไปด้วย เพื่อให้หน้าเว็บเอาไปคำนวณแบบ Real-time
    results = [
        {
            'id': p.id,
            'text': f"{p.code} - {p.name}",
            'cost_price': float(p.cost_price or 0) # ส่งราคาไปให้หน้าเว็บ
        } for p in products
    ]
    return JsonResponse({'results': results})

# ==========================================
# 🌟 [NEW] กระดานสำหรับหัวหน้าแผนกผลิต (ช่าง) 🌟
# ==========================================
@login_required
def production_head_board(request):
    today = timezone.now().date()
    default_start = today - datetime.timedelta(days=30)
    default_end = today + datetime.timedelta(days=30)

    current_monday = today - datetime.timedelta(days=today.weekday())
    current_sunday = current_monday + datetime.timedelta(days=6)
    next_monday = current_monday + datetime.timedelta(days=7)
    next_sunday = next_monday + datetime.timedelta(days=6)

    thai_year_curr = str(current_sunday.year + 543)[-2:]
    thai_year_next = str(next_sunday.year + 543)[-2:]

    if current_monday.month == current_sunday.month:
        str_current_week = f"{current_monday.day}-{current_sunday.day}/{current_monday.month}/{thai_year_curr}"
    else:
        str_current_week = f"{current_monday.day}/{current_monday.month}-{current_sunday.day}/{current_sunday.month}/{thai_year_curr}"

    if next_monday.month == next_sunday.month:
        str_next_week = f"{next_monday.day}-{next_sunday.day}/{next_monday.month}/{thai_year_next}"
    else:
        str_next_week = f"{next_monday.day}/{next_monday.month}-{next_sunday.day}/{next_sunday.month}/{thai_year_next}"

    start_date = request.GET.get('start_date', default_start.strftime('%Y-%m-%d'))
    end_date = request.GET.get('end_date', default_end.strftime('%Y-%m-%d'))
    search_q = request.GET.get('q', '')
    search_branch = request.GET.get('branch', '')
    search_team = request.GET.get('team', '')
    search_salesperson = request.GET.get('salesperson', '')

    orders = ProductionOrder.objects.select_related(
        'product', 'branch', 'production_team', 'salesperson', 'salesperson__branch', 'quotation_ref'
    ).prefetch_related('completed_departments', 'qc_logs').filter(
        # 🌟 [FIXED] เพิ่ม 'WAITING_MATERIALS' และ 'WAITING_INVENTORY' เข้ามาในกระดาน
        status__in=['WAITING_MATERIALS', 'WAITING_INVENTORY', 'PLANNED', 'IN_PROGRESS', 'REWORK'], is_closed=False
    ).order_by('start_date', '-id')

    emp = getattr(request.user, 'employee', None)
    is_manager = False

    if request.user.is_superuser:
        is_manager = True
    elif emp:
        rank = emp.business_rank.lower() if emp.business_rank else ""
        job_title = emp.position.title.lower() if emp.position else ""
        if rank in ['manager', 'director'] or 'manager' in job_title:
            is_manager = True

    if not is_manager and emp and emp.production_team:
        orders = orders.filter(production_team=emp.production_team)

    if start_date: orders = orders.filter(start_date__gte=start_date)
    if end_date: orders = orders.filter(start_date__lte=end_date)
    if search_q:
        orders = orders.filter(Q(code__icontains=search_q) | Q(customer_name__icontains=search_q) | Q(product__name__icontains=search_q))
    if search_branch: orders = orders.filter(branch_id=search_branch)
    if search_team: orders = orders.filter(production_team_id=search_team)
    if search_salesperson: orders = orders.filter(salesperson_id=search_salesperson)

    def format_week_range(d):
        iso_y, iso_w, _ = d.isocalendar()
        monday = datetime.date.fromisocalendar(iso_y, iso_w, 1)
        sunday = monday + datetime.timedelta(days=6)
        y_str = str(sunday.year + 543)[-2:]
        if monday.month == sunday.month:
            return f"จ.{monday.day}-อา.{sunday.day}/{monday.month}/{y_str}"
        else:
            return f"จ.{monday.day}/{monday.month}-อา.{sunday.day}/{sunday.month}/{y_str}"

    col1_upcoming = []
    col2_current = []
    col3_overdue = []
    col4_rework = []

    for order in orders:
        order.display_cohort = format_week_range(order.start_date)

        # งานตีกลับจาก QC ให้แยกไปอยู่คอลัมน์ขวาสุดเสมอ
        if order.status == 'REWORK':
            col4_rework.append(order)
        else:
            # 🌟 [FIXED] จัดงานลงคอลัมน์ตาม "วันที่เริ่มผลิต" ล้วนๆ ไม่ว่าสถานะจะเป็นอะไร
            if order.start_date >= next_monday:
                col1_upcoming.append(order) # งานสัปดาห์หน้า (อนาคต)
            elif order.start_date >= current_monday:
                col2_current.append(order)  # งานสัปดาห์นี้ (ปัจจุบัน)
            else:
                col3_overdue.append(order)  # งานที่เลยกำหนดสัปดาห์ที่แล้ว (อดีต/ค้าง)

    branches = MfgBranch.objects.all().order_by('name')
    teams = ProductionTeam.objects.all().order_by('name')
    salespersons = Salesperson.objects.select_related('branch').all().order_by('name')
    prod_statuses = ProductionStatus.objects.all().order_by('sequence', 'id')

    return render(request, 'manufacturing/production_head_board.html', {
        'orders': orders,
        'col1_upcoming': col1_upcoming,
        'col2_current': col2_current,
        'col3_overdue': col3_overdue,
        'col4_rework': col4_rework,
        'str_current_week': str_current_week,
        'str_next_week': str_next_week,
        'branches': branches,
        'teams': teams,
        'salespersons': salespersons,
        'prod_statuses': prod_statuses,
        'start_date': start_date,
        'end_date': end_date,
        'search_q': search_q,
        'search_branch': search_branch,
        'search_team': search_team,
        'search_salesperson': search_salesperson,
    })

@login_required
def submit_to_qc(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    if order.status in ['IN_PROGRESS', 'REWORK']:
        order.status = 'WAITING_QC'
        order.save()
        messages.success(request, f"✅ ส่งงาน {order.code} ให้แผนก QC ตรวจสอบเรียบร้อยแล้ว!")
    return redirect('production_head_board')

# ==========================================
# 🌟 [NEW] กระดานสำหรับหัวหน้าแผนก QC (ตรวจสอบคุณภาพ) 🌟
# ==========================================
@login_required
def qc_board(request):
    today = timezone.now().date()
    default_start = today - datetime.timedelta(days=30)
    default_end = today + datetime.timedelta(days=30)

    current_monday = today - datetime.timedelta(days=today.weekday())
    current_sunday = current_monday + datetime.timedelta(days=6)
    thai_year_curr = str(current_sunday.year + 543)[-2:]

    if current_monday.month == current_sunday.month:
        str_current_week = f"{current_monday.day}-{current_sunday.day}/{current_monday.month}/{thai_year_curr}"
    else:
        str_current_week = f"{current_monday.day}/{current_monday.month}-{current_sunday.day}/{current_sunday.month}/{thai_year_curr}"

    start_date = request.GET.get('start_date', default_start.strftime('%Y-%m-%d'))
    end_date = request.GET.get('end_date', default_end.strftime('%Y-%m-%d'))
    search_q = request.GET.get('q', '')
    search_branch = request.GET.get('branch', '')
    search_team = request.GET.get('team', '')
    search_salesperson = request.GET.get('salesperson', '')

    orders = ProductionOrder.objects.select_related(
        'product', 'branch', 'production_team', 'salesperson', 'salesperson__branch', 'quotation_ref'
    ).prefetch_related('qc_logs').filter(
        status__in=['WAITING_QC', 'REWORK'], is_closed=False
    ).order_by('start_date', '-id')

    if start_date: orders = orders.filter(start_date__gte=start_date)
    if end_date: orders = orders.filter(start_date__lte=end_date)
    if search_q:
        orders = orders.filter(Q(code__icontains=search_q) | Q(customer_name__icontains=search_q) | Q(product__name__icontains=search_q))
    if search_branch: orders = orders.filter(branch_id=search_branch)
    if search_team: orders = orders.filter(production_team_id=search_team)
    if search_salesperson: orders = orders.filter(salesperson_id=search_salesperson)

    def format_week_range(d):
        iso_y, iso_w, _ = d.isocalendar()
        monday = datetime.date.fromisocalendar(iso_y, iso_w, 1)
        sunday = monday + datetime.timedelta(days=6)
        y_str = str(sunday.year + 543)[-2:]
        if monday.month == sunday.month:
            return f"จ.{monday.day}-อา.{sunday.day}/{monday.month}/{y_str}"
        else:
            return f"จ.{monday.day}/{monday.month}-อา.{sunday.day}/{sunday.month}/{y_str}"

    col1_current = []
    col2_overdue = []
    col3_rework = []

    for order in orders:
        order.display_cohort = format_week_range(order.start_date)

        if order.status == 'REWORK':
            col3_rework.append(order)
        elif order.status == 'WAITING_QC':
            if order.start_date < current_monday:
                col2_overdue.append(order)
            else:
                col1_current.append(order)

    branches = MfgBranch.objects.all().order_by('name')
    teams = ProductionTeam.objects.all().order_by('name')
    salespersons = Salesperson.objects.select_related('branch').all().order_by('name')

    return render(request, 'manufacturing/qc_board.html', {
        'orders': orders,
        'col1_current': col1_current,
        'col2_overdue': col2_overdue,
        'col3_rework': col3_rework,
        'str_current_week': str_current_week,
        'branches': branches,
        'teams': teams,
        'salespersons': salespersons,
        'start_date': start_date,
        'end_date': end_date,
        'search_q': search_q,
        'search_branch': search_branch,
        'search_team': search_team,
        'search_salesperson': search_salesperson,
    })

@login_required
@transaction.atomic
def process_qc(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)

    if request.method == 'POST':
        action = request.POST.get('action')

        order.qc_paint = request.POST.get('qc_paint') == 'on'
        order.qc_internal = request.POST.get('qc_internal') == 'on'
        order.qc_external = request.POST.get('qc_external') == 'on'
        order.qc_electrical = request.POST.get('qc_electrical') == 'on'
        order.qc_plumbing = request.POST.get('qc_plumbing') == 'on'
        order.qc_aircon = request.POST.get('qc_aircon') == 'on'

        if action == 'save_progress':
            comments = request.POST.get('comments', '')

            # 1. กวาดไฟล์ "ผ่าน" บันทึกลงตารางทันที
            passed_files = request.FILES.getlist('passed_media')
            for f in passed_files:
                QCPassedMedia.objects.create(
                    production_order=order,
                    uploaded_by=getattr(request.user, 'employee', None),
                    file=f
                )

            # 2. กวาดไฟล์ "ไม่ผ่าน" และคอมเมนต์ บันทึกเป็นร่าง(Draft) ในรอบปัจจุบัน
            failed_files = request.FILES.getlist('failed_media')
            if failed_files or comments:
                log, created = QCInspectionLog.objects.get_or_create(
                    production_order=order,
                    round_number=order.rework_count + 1,
                    defaults={
                        'inspector': getattr(request.user, 'employee', None),
                        'status': 'FAILED',
                        'comments': comments
                    }
                )
                if not created and comments:
                    log.comments = comments
                    log.save()

                for f in failed_files:
                    QCFailedMedia.objects.create(qc_log=log, file=f)

            order.save()
            messages.success(request, f"💾 บันทึกความคืบหน้าและรูปภาพสำหรับ {order.code} ลงฐานข้อมูลเรียบร้อยแล้ว!")
            return redirect('qc_board')

        elif action == 'pass':
            alloc_code = f"{order.product.code}-{order.code}"
            alloc_name = f"{order.product.name} [{order.code}]"
            if order.customer_name: alloc_name += f" (ลค. {order.customer_name})"

            allocated_product, created = Product.objects.get_or_create(
                code=alloc_code,
                defaults={
                    'name': alloc_name[:255], 'category': order.product.category,
                    'product_type': 'FG', 'sell_price': order.product.sell_price,
                    'cost_price': order.product.cost_price, 'min_level': 0,
                    'stock_qty': 0, 'is_active': True,
                    'image': order.product.image if order.product.image else None,
                }
            )

            ref_doc_text = ""
            if order.quotation_ref:
                ref_doc_text = f" อ้างอิงใบเสนอราคา: {order.quotation_ref.code}"

            note_text = f"รับสินค้าจากฝ่ายผลิต {order.code} (ผ่าน QC){ref_doc_text}"

            qc_doc = QCReceiptDoc.objects.create(production_order=order, product=allocated_product, quantity=Decimal(str(order.quantity)), created_by=request.user)
            StockMovement.objects.create(doc=None, reference_doc=qc_doc.doc_no, note=note_text, product=allocated_product, quantity=Decimal(str(order.quantity)), movement_type='IN', created_by=request.user)

            order.status = 'COMPLETED'
            order.finish_date = timezone.now().date()
            order.is_qc_passed = True
            order.save()

            # 🌟 [NEW] บันทึกประวัติการตรวจ "ผ่าน" ลงในฐานข้อมูล เพื่อให้โชว์ในหน้า My QC Log
            QCInspectionLog.objects.create(
                production_order=order,
                round_number=order.rework_count + 1,
                inspector=getattr(request.user, 'employee', None),
                status='PASSED',
                comments="ผ่านการตรวจสอบเรียบร้อย"
            )

            passed_files = request.FILES.getlist('passed_media')
            for f in passed_files:
                QCPassedMedia.objects.create(
                    production_order=order,
                    uploaded_by=getattr(request.user, 'employee', None),
                    file=f
                )

            messages.success(request, f"🎉 ตรวจผ่านเรียบร้อย! รับ {allocated_product.name} เข้าคลังแล้ว พร้อมบันทึกคลังภาพ/คลิป")

        elif action == 'fail':
            comments = request.POST.get('comments', '')
            order.rework_count += 1
            order.status = 'REWORK'

            # ดึง Log เดิมที่อาจจะถูกร่างไว้ตอนกด "บันทึกไว้ก่อน" ขึ้นมาใช้งานจริง
            log, created = QCInspectionLog.objects.get_or_create(
                production_order=order,
                round_number=order.rework_count,
                defaults={
                    'inspector': getattr(request.user, 'employee', None),
                    'status': 'FAILED',
                    'comments': comments
                }
            )

            log.comments = comments
            log.inspector = getattr(request.user, 'employee', None)

            if 'defect_image_1' in request.FILES:
                log.defect_image_1 = request.FILES['defect_image_1']
            if 'defect_image_2' in request.FILES:
                log.defect_image_2 = request.FILES['defect_image_2']
            log.save()

            failed_files = request.FILES.getlist('failed_media')
            for f in failed_files:
                QCFailedMedia.objects.create(
                    qc_log=log,
                    file=f
                )

            order.save()
            messages.warning(request, f"🚨 ตีกลับงาน {order.code} ให้ช่างแก้ไขเรียบร้อย พร้อมแนบรูป/คลิปตำหนิ (รอบที่ {order.rework_count})")

    return redirect('qc_board')

# ==========================================
# 🌟 [NEW] ศูนย์ตรวจสอบแบบแปลนและตั้งเบิก (Blueprint Hub) 🌟
# ==========================================
from .models import BlueprintClaim, BlueprintLog

@login_required
def blueprint_hub(request):
    current_emp = getattr(request.user, 'employee', None)

    date_range = request.GET.get('date_range', '')
    start_date = ''
    end_date = ''

    if date_range:
        parts = date_range.split(' - ')
        if len(parts) == 2:
            try:
                d1, m1, y1 = parts[0].strip().split('/')
                start_date = f"{y1}-{m1}-{d1}"
                d2, m2, y2 = parts[1].strip().split('/')
                end_date = f"{y2}-{m2}-{d2}"
            except:
                pass

    q_job = request.GET.get('q_job', '')
    q_customer = request.GET.get('q_customer', '')
    search_sp = request.GET.get('salesperson', '')
    claim_status = request.GET.get('claim_status', '')

    pending_jobs = ProductionOrder.objects.select_related('product', 'salesperson', 'salesperson__branch').filter(
        status__in=['WAITING_BLUEPRINT', 'PLANNED'],
        is_closed=False
    ).order_by('start_date')

    # 🌟 [FIXED] เปลี่ยนจาก blueprint_approved_by เป็น blueprint_approved_at__isnull=False (แอดมินกดอนุมัติก็จะโชว์)
    history_jobs = ProductionOrder.objects.select_related(
        'product', 'salesperson', 'salesperson__branch', 'blueprint_approved_by', 'blueprint_claim'
    ).filter(blueprint_approved_at__isnull=False)

    if start_date:
        history_jobs = history_jobs.filter(blueprint_approved_at__date__gte=start_date)
    if end_date:
        history_jobs = history_jobs.filter(blueprint_approved_at__date__lte=end_date)
    if q_job:
        history_jobs = history_jobs.filter(code__icontains=q_job)
    if q_customer:
        history_jobs = history_jobs.filter(customer_name__icontains=q_customer)
    if search_sp:
        history_jobs = history_jobs.filter(salesperson_id=search_sp)

    if claim_status == 'UNCLAIMED':
        history_jobs = history_jobs.filter(blueprint_claim__isnull=True)
    elif claim_status in ['PENDING', 'PAID', 'REJECTED']:
        history_jobs = history_jobs.filter(blueprint_claim__status=claim_status)

    history_jobs = history_jobs.order_by('-blueprint_approved_at')[:100]

    # 🌟 [FIXED] เพิ่ม select_related เพื่อให้ระบบดึงข้อมูลเซลส์ ใบเสนอราคา และใบขาย มารอไว้เลย ป้องกันเว็บโหลดช้า
    claimable_jobs = ProductionOrder.objects.select_related(
        'salesperson', 'quotation_ref', 'quotation_ref__invoice'
    ).filter(
        blueprint_approved_by=current_emp,
        blueprint_claim__isnull=True,
        quotation_ref__invoice__isnull=False
    ).order_by('blueprint_approved_at')

    waiting_invoice_jobs = ProductionOrder.objects.select_related(
        'salesperson', 'quotation_ref'
    ).filter(
        blueprint_approved_by=current_emp,
        blueprint_claim__isnull=True,
        quotation_ref__invoice__isnull=True
    ).order_by('blueprint_approved_at')

    claims = BlueprintClaim.objects.filter(employee=current_emp).order_by('-created_at')

    salespersons = Salesperson.objects.select_related('branch').all().order_by('name')

    return render(request, 'manufacturing/blueprint_hub.html', {
        'pending_jobs': pending_jobs,
        'history_jobs': history_jobs,
        'claimable_jobs': claimable_jobs,
        'waiting_invoice_jobs': waiting_invoice_jobs,
        'claims': claims,
        'salespersons': salespersons,
        'date_range': date_range,
        'q_job': q_job,
        'q_customer': q_customer,
        'search_sp': search_sp,
        'claim_status': claim_status,
    })

@login_required
def blueprint_workspace(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    materials = order.materials.all()
    has_bom = BOM.objects.filter(product=order.product).exists()
    raw_materials = Product.objects.filter(product_type='RM', is_active=True).order_by('code')
    logs = order.blueprint_logs.all()

    return render(request, 'manufacturing/blueprint_workspace.html', {
        'order': order,
        'materials': materials,
        'has_bom': has_bom,
        'raw_materials': raw_materials,
        'logs': logs
    })

@login_required
def blueprint_approve(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    current_emp = getattr(request.user, 'employee', None)

    if request.method == 'POST':
        order.status = 'WAITING_MATERIALS'
        order.blueprint_approved_by = current_emp
        order.blueprint_approved_at = timezone.now()
        order.save()

        BlueprintLog.objects.create(
            production_order=order,
            action="ตรวจสอบความถูกต้องของแบบแปลนและโครงสร้างวัตถุดิบ (Approve)",
            employee=current_emp
        )

        messages.success(request, f"✅ ยืนยันความถูกต้องของ JOB {order.code} เรียบร้อย! งานถูกส่งต่อไปยังแผนกสั่งวัตถุดิบแล้วค่ะ")
        # 🌟 [FIXED] สั่งให้รีไดเรกต์แล้วกระโดดไปที่แท็บประวัติอัตโนมัติ
        return redirect(reverse('blueprint_hub') + '#history')

    return redirect('blueprint_workspace', pk=pk)

@login_required
def blueprint_create_claim(request):
    if request.method == 'POST':
        job_ids = request.POST.getlist('job_ids')
        current_emp = getattr(request.user, 'employee', None)

        if job_ids and current_emp:
            group = current_emp.sales_group
            if not group or group.flat_rate_amount <= 0:
                messages.error(request, "❌ ไม่สามารถตั้งเบิกได้: คุณยังไม่มีสังกัดทีม หรือ ทีมยังไม่ได้ตั้งค่า 'ค่าตอบแทนเหมาจ่าย'")
                return redirect('blueprint_hub')

            jobs = ProductionOrder.objects.filter(id__in=job_ids, blueprint_approved_by=current_emp, blueprint_claim__isnull=True)
            if jobs.exists():
                rate_per_job = Decimal(str(group.flat_rate_amount))
                total_jobs = jobs.count()
                total_amount = Decimal(str(total_jobs)) * rate_per_job

                claim = BlueprintClaim.objects.create(employee=current_emp, total_jobs=total_jobs, total_amount=total_amount)
                jobs.update(blueprint_claim=claim)

                members = group.members.all()
                l1_count = members.filter(group_role='LEVEL1').count() or 1
                l2_count = members.filter(group_role='LEVEL2').count() or 1

                for member in members:
                    share_pct = Decimal('0')
                    if member.group_role == 'LEADER':
                        share_pct = Decimal(str(group.share_leader))
                    elif member.group_role == 'LEVEL1':
                        share_pct = Decimal(str(group.share_level1)) / Decimal(str(l1_count))
                    elif member.group_role == 'LEVEL2':
                        share_pct = Decimal(str(group.share_level2)) / Decimal(str(l2_count))

                    if share_pct > 0:
                        split_amount = (total_amount * share_pct) / Decimal('100')
                        BlueprintClaimSplit.objects.create(
                            claim=claim, employee=member, role_name=member.get_group_role_display(),
                            percentage=share_pct, amount=split_amount
                        )

                messages.success(request, f"💰 สร้างใบตั้งเบิก {claim.code} สำเร็จ! และระบบกระจายยอดเงินให้ลูกทีมอัตโนมัติเรียบร้อยค่ะ")
            else:
                messages.error(request, "❌ ไม่พบงานที่สามารถตั้งเบิกได้ หรือมีการเบิกไปแล้ว")
        else:
            messages.warning(request, "⚠️ กรุณาเลือกงานอย่างน้อย 1 รายการ")
    return redirect('blueprint_hub')

def get_thai_baht_text(number):
    if number == 0: return "ศูนย์บาทถ้วน"
    import math
    number = round(number, 2)
    baht = math.floor(number)
    satang = int((number - baht) * 100)

    def read_num(n):
        if n == 0: return ""
        numbers = ["", "หนึ่ง", "สอง", "สาม", "สี่", "ห้า", "หก", "เจ็ด", "แปด", "เก้า"]
        positions = ["", "สิบ", "ร้อย", "พัน", "หมื่น", "แสน", "ล้าน"]
        s = str(n)
        length = len(s)
        res = ""
        for i, digit in enumerate(s):
            val = int(digit)
            pos = length - i - 1
            if val == 0:
                continue
            if pos == 0 and val == 1 and length > 1:
                res += "เอ็ด"
            elif pos == 1 and val == 1:
                res += "สิบ"
            elif pos == 1 and val == 2:
                res += "ยี่สิบ"
            else:
                res += numbers[val] + positions[pos]
        return res

    res = ""
    if baht > 0:
        res += read_num(baht) + "บาท"
    if satang > 0:
        res += read_num(satang) + "สตางค์"
    else:
        res += "ถ้วน"
    return res

@login_required
def print_blueprint_claim(request, pk):
    claim = get_object_or_404(BlueprintClaim, pk=pk)
    company = CompanyInfo.objects.first()
    jobs = claim.production_orders.all()
    amount_text = get_thai_baht_text(claim.total_amount)

    splits = claim.splits.all()

    total_split_amount = sum([s.amount for s in splits]) if splits else Decimal('0')
    fund_amount = claim.total_amount - total_split_amount

    total_split_percent = sum([s.percentage for s in splits]) if splits else Decimal('0')
    fund_percentage = Decimal('100') - total_split_percent

    return render(request, 'manufacturing/print_blueprint_claim.html', {
        'claim': claim,
        'jobs': jobs,
        'company': company,
        'amount_text': amount_text,
        'fund_amount': fund_amount,
        'fund_percentage': fund_percentage,
    })

# ==========================================
# 🌟 [NEW] กระดานสำหรับแผนกจัดส่งสินค้า (Logistics Board) 🌟
# ==========================================
@login_required
def logistics_board(request):
    orders = ProductionOrder.objects.select_related(
        'product', 'branch', 'transporter', 'delivery_status', 'quotation_ref', 'logistics_claim'
    ).filter(
        Q(status='COMPLETED') | Q(status='WAITING_QC') | Q(is_onsite=True)
    ).order_by('delivery_date', '-id')

    col1_pending, col2_delivering, col3_delivered = [], [], []

    for order in orders:
        # 🌟 [NEW] ลอจิกตรวจสอบยอดค้างชำระที่แท้จริงจาก Invoice (ถ้ามี) 🌟
        order.actual_balance = Decimal('0')
        if order.quotation_ref:
            inv = None
            if hasattr(order.quotation_ref, 'invoice_set'):
                inv = order.quotation_ref.invoice_set.first()
            elif hasattr(order.quotation_ref, 'invoice'):
                inv = order.quotation_ref.invoice

            if inv:
                # ถ้ามีการเปิด Invoice แล้ว ยึดยอดคงเหลือจาก Invoice เป็นหลัก
                order.actual_balance = getattr(inv, 'balance_amount', Decimal('0'))
            else:
                # ถ้ายังไม่เปิด Invoice ให้ดูยอดคงค้างจาก Quotation
                grand_total = getattr(order.quotation_ref, 'grand_total', Decimal('0'))
                deposit = getattr(order.quotation_ref, 'deposit_amount', Decimal('0'))
                order.actual_balance = getattr(order.quotation_ref, 'balance_due', grand_total - deposit)

        # จัดกลุ่มลงคอลัมน์
        if order.transporter is None:
            col1_pending.append(order)
        elif order.delivery_status and order.delivery_status.name in ['ส่งมอบสำเร็จ', 'ลูกค้าเซ็นรับแล้ว', 'จัดส่งเรียบร้อย']:
            col3_delivered.append(order)
        else:
             col2_delivering.append(order)

    transporters = Transporter.objects.all().order_by('name')
    delivery_statuses = DeliveryStatus.objects.all().order_by('name')
    claims = LogisticsClaim.objects.all().order_by('-created_at')[:20]

    return render(request, 'manufacturing/logistics_board.html', {
        'orders': orders, 'col1_pending': col1_pending, 'col2_delivering': col2_delivering,
        'col3_delivered': col3_delivered, 'transporters': transporters,
        'delivery_statuses': delivery_statuses, 'claims': claims
    })

@login_required
@transaction.atomic
def process_logistics(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'assign_truck':
            transporter_id = request.POST.get('transporter')
            delivery_fee = request.POST.get('delivery_fee', 0)
            delivery_date_str = request.POST.get('delivery_date')

            if transporter_id:
                order.transporter_id = transporter_id

                if delivery_date_str:
                    order.delivery_date = parse_date(delivery_date_str)

                try:
                    order.delivery_fee = Decimal(str(delivery_fee).replace(',', ''))
                except:
                    order.delivery_fee = Decimal('0.00')

                messages.success(request, f"🚛 จ่ายงาน {order.code} ให้ทีมขนส่งและระบุค่าจ้างเรียบร้อยแล้ว!")

        elif action == 'update_status':
            delivery_status_id = request.POST.get('delivery_status')
            if delivery_status_id:
                status_obj = DeliveryStatus.objects.get(id=delivery_status_id)

                # ตรวจสอบว่าเป็นการส่งมอบสำเร็จหรือไม่
                is_delivered_success = status_obj.name in ['ส่งมอบสำเร็จ', 'ลูกค้าเซ็นรับแล้ว', 'จัดส่งเรียบร้อย']

                # ถ้าเป็นสถานะส่งมอบสำเร็จ
                if is_delivered_success:
                    # 🌟 เช็คว่ามีรูปลายเซ็นในระบบแล้ว หรือกำลังอัปโหลดเข้ามาใหม่
                    has_proof = order.proof_of_delivery or 'proof_of_delivery' in request.FILES

                    if has_proof:
                        if 'proof_of_delivery' in request.FILES:
                            order.proof_of_delivery = request.FILES['proof_of_delivery']

                        order.delivery_status_id = delivery_status_id
                        # 🌟 [DELETED] นำระบบตัดสต็อกหน้าลานตรงนี้ออกไปแล้ว ตามคำขอของคุณปอนด์
                        messages.success(request, "📦 อัปเดตสถานะการส่งมอบสำเร็จ!")
                    else:
                        messages.error(request, "❌ กรุณาแนบรูปภาพใบส่งมอบสินค้า (ที่มีลายเซ็นลูกค้า) ก่อนปิดงาน!")
                        return redirect('logistics_board')
                else:
                    order.delivery_status_id = delivery_status_id
                    messages.success(request, f"📦 อัปเดตสถานะการจัดส่ง {order.code} เรียบร้อยแล้ว!")
        order.save()
    return redirect('logistics_board')

@login_required
def create_logistics_claim(request):
    if request.method == 'POST':
        job_ids = request.POST.getlist('job_ids')
        if job_ids:
            jobs = ProductionOrder.objects.filter(id__in=job_ids, logistics_claim__isnull=True)
            if jobs.exists():
                transporter = jobs.first().transporter
                total_amt = sum([j.delivery_fee for j in jobs])
                claim = LogisticsClaim.objects.create(transporter=transporter, total_jobs=jobs.count(), total_amount=total_amt)
                jobs.update(logistics_claim=claim)
                messages.success(request, f"💰 สร้างใบตั้งเบิกค่ารถ {claim.code} สำเร็จ!")
            else: messages.error(request, "❌ ไม่พบงาน หรือมีการตั้งเบิกไปแล้ว")
        else: messages.warning(request, "⚠️ กรุณาติ๊กเลือกอย่างน้อย 1 งาน")
    return redirect('logistics_claim_history')

@login_required
def print_delivery_note(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    company = CompanyInfo.objects.first()

    amount_text = ""
    if order.quotation_ref and order.quotation_ref.grand_total:
        amount_text = get_thai_baht_text(order.quotation_ref.grand_total)

    return render(request, 'manufacturing/print_delivery_note.html', {
        'order': order,
        'company': company,
        'amount_text': amount_text
    })

@login_required
def print_consent_form(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    company = CompanyInfo.objects.first()
    return render(request, 'manufacturing/print_consent_form.html', {
        'order': order,
        'company': company,
    })

# 🌟 [NEW] ฟังก์ชันดึงข้อมูลเพื่อพิมพ์ใบตรวจสอบก่อนติดตั้ง 🌟
@login_required
def print_site_checklist(request, pk):
    order = get_object_or_404(ProductionOrder, pk=pk)
    company = CompanyInfo.objects.first()
    return render(request, 'manufacturing/print_site_checklist.html', {
        'order': order,
        'company': company,
    })

@login_required
def print_logistics_claim(request, pk):
    claim = get_object_or_404(LogisticsClaim, pk=pk)
    company = CompanyInfo.objects.first()
    jobs = claim.production_orders.all()
    amount_text = get_thai_baht_text(claim.total_amount)
    return render(request, 'manufacturing/print_logistics_claim.html', {'claim': claim, 'jobs': jobs, 'company': company, 'amount_text': amount_text})

# ==========================================
# 🌟 [NEW] หน้าต่างรวมประวัติการจัดส่งสินค้า (Master Logistics Log) 🌟
# ==========================================
@login_required
def logistics_history_log(request):
    # ดึงเฉพาะงานที่ถูกจัดส่งเรียบร้อยแล้ว หรือเสร็จสมบูรณ์แล้ว
    orders = ProductionOrder.objects.select_related(
        'product', 'branch', 'transporter', 'delivery_status', 'salesperson'
    ).filter(
        Q(delivery_status__name__in=['ส่งมอบสำเร็จ', 'ลูกค้าเซ็นรับแล้ว', 'จัดส่งเรียบร้อย']) |
        Q(is_onsite=True, status='COMPLETED')
    ).order_by('-delivery_date', '-finish_date', '-id')

    search_q = request.GET.get('q', '')
    search_transporter = request.GET.get('transporter', '')
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')

    if search_q:
        orders = orders.filter(
            Q(code__icontains=search_q) |
            Q(customer_name__icontains=search_q) |
            Q(product__name__icontains=search_q)
        )
    if search_transporter:
        orders = orders.filter(transporter_id=search_transporter)
    if start_date:
        orders = orders.filter(delivery_date__gte=start_date)
    if end_date:
        orders = orders.filter(delivery_date__lte=end_date)

    paginator = Paginator(orders, 30) # โชว์หน้าละ 30 รายการ
    page_num = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_num)

    transporters = Transporter.objects.all().order_by('name')

    return render(request, 'manufacturing/logistics_history_log.html', {
        'page_obj': page_obj,
        'search_q': search_q,
        'search_transporter': search_transporter,
        'start_date': start_date,
        'end_date': end_date,
        'transporters': transporters,
    })

@login_required
def logistics_claim_history(request):
    claims = LogisticsClaim.objects.select_related('transporter').prefetch_related('production_orders', 'production_orders__quotation_ref').all().order_by('-created_at')

    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')
    q_transporter = request.GET.get('q_transporter', '')
    q_job = request.GET.get('q_job', '')
    q_customer = request.GET.get('q_customer', '')

    if start_date: claims = claims.filter(created_at__date__gte=start_date)
    if end_date: claims = claims.filter(created_at__date__lte=end_date)
    if q_transporter: claims = claims.filter(Q(transporter__name__icontains=q_transporter) | Q(transporter__driver_name__icontains=q_transporter))
    if q_job: claims = claims.filter(production_orders__code__icontains=q_job).distinct()
    if q_customer: claims = claims.filter(production_orders__customer_name__icontains=q_customer).distinct()

    paginator = Paginator(claims, 20)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'manufacturing/logistics_claim_history.html', {
        'page_obj': page_obj, 'start_date': start_date, 'end_date': end_date,
        'q_transporter': q_transporter, 'q_job': q_job, 'q_customer': q_customer
    })

@login_required
def ajax_add_transporter_full(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        driver_name = request.POST.get('driver_name', '')
        vehicle_plate = request.POST.get('vehicle_plate', '')
        address = request.POST.get('address', '')
        bank_account = request.POST.get('bank_account', '')
        id_card_image = request.FILES.get('id_card_image')

        if name:
            obj, created = Transporter.objects.get_or_create(name=name, defaults={
                'driver_name': driver_name, 'vehicle_plate': vehicle_plate,
                'address': address, 'bank_account': bank_account, 'id_card_image': id_card_image
            })
            return JsonResponse({'success': True, 'id': obj.id, 'name': obj.name})
    return JsonResponse({'success': False})

@login_required
@transaction.atomic
def pay_logistics_claim(request, claim_id):
    claim = get_object_or_404(LogisticsClaim, pk=claim_id)

    if request.method == 'POST' and claim.status == 'PENDING':
        claim.status = 'PAID'
        claim.save()

        Expense.objects.create(
            title=f"จ่ายค่าขนส่ง/ทีมรถ (ใบเบิก: {claim.code}) - {claim.transporter.name}",
            amount=claim.total_amount,
            date=timezone.now().date(),
        )

        messages.success(request, f"✅ ทำจ่ายเงินใบเบิก {claim.code} จำนวน ฿{claim.total_amount:,.2f} สำเร็จ! และระบบบันทึกลงบัญชีรายจ่ายให้อัตโนมัติแล้วค่ะ")

    return redirect('logistics_claim_history')

@login_required
@transaction.atomic
def pay_blueprint_claim(request, claim_id):
    claim = get_object_or_404(BlueprintClaim, pk=claim_id)

    if request.method == 'POST' and claim.status == 'PENDING':
        claim.status = 'PAID'
        claim.save()

        emp_name = claim.employee.first_name if claim.employee else 'ไม่ระบุ'
        Expense.objects.create(
            title=f"จ่ายค่าแบบแปลน/ทีมช่าง (ใบเบิก: {claim.code}) - {emp_name}",
            amount=claim.total_amount,
            date=timezone.now().date(),
        )

        messages.success(request, f"✅ ทำจ่ายเงินใบเบิก {claim.code} จำนวน ฿{claim.total_amount:,.2f} สำเร็จ! และระบบบันทึกลงบัญชีรายจ่ายให้อัตโนมัติแล้วค่ะ")

    return redirect('blueprint_hub')

@login_required
def update_bom_labor_cost(request, pk):
    if request.method == 'POST':
        bom = get_object_or_404(BOM, pk=pk)
        try:
            # 1. ค่าแรงประกอบ
            labor_val = request.POST.get('labor_cost', '0').replace(',', '')
            bom.labor_cost = Decimal(labor_val)

            # 2. ค่าบริหารการขาย
            bom.sales_admin_type = request.POST.get('sales_admin_type', 'PERCENT')
            bom.sales_admin_value = Decimal(request.POST.get('sales_admin_value', '0').replace(',', ''))

            # 3. ค่าขนส่ง
            bom.freight_type = request.POST.get('freight_type', 'THB')
            bom.freight_value = Decimal(request.POST.get('freight_value', '0').replace(',', ''))

            # 4. ค่าใช้จ่ายอื่นๆ
            bom.misc_expense_type = request.POST.get('misc_expense_type', 'THB')
            bom.misc_expense_value = Decimal(request.POST.get('misc_expense_value', '0').replace(',', ''))

            bom.save()
            messages.success(request, f"✅ บันทึกต้นทุนการผลิตสำหรับ {bom.product.code} เรียบร้อยแล้ว")
        except (ValueError, InvalidOperation):
            messages.error(request, "❌ เกิดข้อผิดพลาดในการบันทึกตัวเลขต้นทุน")
    return redirect('bom_list')

@login_required
def bom_list(request):
    boms = BOM.objects.select_related('product').annotate(
        item_count=Count('items', distinct=True),
        total_rm_cost=Coalesce(
            Sum(F('items__quantity') * F('items__raw_material__cost_price')),
            0,
            output_field=DecimalField(max_digits=12, decimal_places=2)
        )
    ).order_by('-id')

    # 🌟 สมองกลคำนวณต้นทุนใหม่แบบละเอียด 🌟
    for bom in boms:
        sell_price = bom.product.sell_price or Decimal('0.00')
        total_rm = bom.total_rm_cost or Decimal('0.00')
        labor = bom.labor_cost or Decimal('0.00')

        # 1. ค่าบริหารการขาย (ถ้าระบุเป็น % ให้หักออก 10% จากราคาขายก่อนคูณ)
        sales_admin = Decimal('0.00')
        if bom.sales_admin_type == 'PERCENT':
            base_for_admin = sell_price * Decimal('0.90')
            sales_admin = base_for_admin * (bom.sales_admin_value / Decimal('100'))
        else:
            sales_admin = bom.sales_admin_value

        # 2. ค่าขนส่ง (ถ้าระบุเป็น % ให้คูณจากราคาขายเต็ม)
        freight = Decimal('0.00')
        if bom.freight_type == 'PERCENT':
            freight = sell_price * (bom.freight_value / Decimal('100'))
        else:
            freight = bom.freight_value

        # 3. ค่าใช้จ่ายอื่นๆ (ถ้าระบุเป็น % ให้คูณจากราคาขายเต็ม)
        misc = Decimal('0.00')
        if bom.misc_expense_type == 'PERCENT':
            misc = sell_price * (bom.misc_expense_value / Decimal('100'))
        else:
            misc = bom.misc_expense_value

        # คำนวณกำไรสุทธิ: ราคาขาย - (ต้นทุนวัตถุดิบ + ค่าแรง + ค่าบริหาร + ค่าขนส่ง + อื่นๆ)
        bom.estimated_profit = sell_price - (total_rm + labor + sales_admin + freight + misc)

        # แนบต้นทุนแต่ละก้อนไปแสดงผลหน้าเว็บด้วย (สำหรับแนวทาง B ที่ต้องแสดงทุกคอลัมน์)
        bom.calculated_sales_admin = sales_admin
        bom.calculated_freight = freight
        bom.calculated_misc = misc

        # คำนวณ % กำไร
        if sell_price > 0:
            bom.profit_percentage = (bom.estimated_profit / sell_price) * Decimal('100')
        else:
            bom.profit_percentage = Decimal('0.00')

    return render(request, 'manufacturing/bom_list.html', {'boms': boms})

# ==========================================
# 🌟 ศูนย์รวมแฟ้มภาพและบริการ (Media Hub)
# ==========================================
@login_required
def media_hub(request, pk):
    # 1. ดึงข้อมูลใบสั่งผลิต (JOB)
    order = get_object_or_404(ProductionOrder, pk=pk)

    # 2. ดึงรูปภาพ/คลิป ที่ผ่าน QC แล้ว (สีเขียว)
    passed_medias = order.passed_media.all().order_by('-id')

    # 3. ดึงประวัติการตีกลับและรูปตำหนิ (สีแดง)
    qc_logs = order.qc_logs.prefetch_related('failed_media').order_by('-round_number')

    # 🌟 [NEW] 4. ดึงประวัติการแจ้งซ่อม (สีเหลือง)
    service_tickets = order.service_tickets.all().order_by('-created_at')

    return render(request, 'manufacturing/media_hub.html', {
        'order': order,
        'passed_medias': passed_medias,
        'qc_logs': qc_logs,
        'service_tickets': service_tickets, # 🌟 ส่งตัวแปรนี้ไปที่หน้าเว็บ
    })

@login_required
def media_hub_list(request):
    search_q = request.GET.get('q', '')

    # ดึงข้อมูลใบสั่งผลิตทั้งหมดเรียงตามวันที่ล่าสุด
    orders = ProductionOrder.objects.select_related(
        'product', 'salesperson', 'production_team'
    ).all().order_by('-id')

    # ระบบค้นหา
    if search_q:
        orders = orders.filter(
            Q(code__icontains=search_q) |
            Q(customer_name__icontains=search_q) |
            Q(product__name__icontains=search_q)
        )

    return render(request, 'manufacturing/media_hub_list.html', {
        'orders': orders[:100],  # ดึงมาแสดงผลเบื้องต้น 100 รายการเพื่อไม่ให้เว็บโหลดช้า
        'search_q': search_q,
    })

# ==========================================
# 🌟 [UPDATED] ประวัติการทำงานส่วนบุคคล และ ศูนย์รวมประวัติสำหรับผู้บริหาร 🌟
# ==========================================
@login_required
def my_production_log(request):
    emp = getattr(request.user, 'employee', None)
    team = emp.production_team if emp else None

    # เช็คสิทธิ์ว่าเป็นผู้บริหารหรือแอดมินหรือไม่
    is_manager = False
    if request.user.is_superuser:
        is_manager = True
    elif emp:
        rank = emp.business_rank.lower() if emp.business_rank else ""
        if rank in ['manager', 'director', 'executive']:
            is_manager = True

    # 🌟 1. กรณีเป็นผู้บริหาร (Admin/Manager): ดึงประวัติที่เสร็จสิ้นของ "ทุกแผนก"
    if is_manager:
        orders = ProductionOrder.objects.select_related(
            'product', 'salesperson', 'production_team', 'transporter', 'blueprint_approved_by'
        ).filter(
            # ดึงงานที่เสร็จสิ้นแล้วในมิติใดมิติหนึ่ง (ประกอบเสร็จ, ผ่าน QC, หรือส่งมอบแล้ว)
            Q(status__in=['COMPLETED', 'CLOSED']) |
            Q(is_qc_passed=True) |
            Q(delivery_status__name__in=['ส่งมอบสำเร็จ', 'ลูกค้าเซ็นรับแล้ว', 'จัดส่งเรียบร้อย'])
        ).order_by('-finish_date', '-id')[:100]

    # 🌟 2. กรณีเป็นช่างทั่วไป: ดึงเฉพาะงานที่ "ทีมตัวเอง" ประกอบเสร็จ
    elif team:
        orders = ProductionOrder.objects.select_related(
            'product', 'salesperson', 'production_team'
        ).filter(
            production_team=team,
            status__in=['COMPLETED', 'CLOSED']
        ).order_by('-finish_date', '-id')[:100]

    # 🌟 3. กรณีไม่มีสังกัดทีมและไม่ใช่ผู้บริหาร
    else:
        orders = ProductionOrder.objects.none()

    return render(request, 'manufacturing/my_production_log.html', {
        'orders': orders,
        'is_manager': is_manager, # ส่งตัวแปรไปบอกหน้าเว็บว่าเป็นผู้บริหาร
    })

@login_required
def my_qc_log(request):
    emp = getattr(request.user, 'employee', None)

    # 🌟 1. เช็คสิทธิ์ว่าเป็นผู้บริหารหรือแอดมินหรือไม่
    is_manager = False
    if request.user.is_superuser:
        is_manager = True
    elif emp:
        rank = emp.business_rank.lower() if emp.business_rank else ""
        if rank in ['manager', 'director', 'executive']:
            is_manager = True

    # 🌟 2. ดึงประวัติแบบแยกสิทธิ์
    if is_manager:
        # แอดมิน/ผู้บริหาร: เห็นประวัติการตรวจของ "ทุกคน"
        qc_logs = QCInspectionLog.objects.select_related(
            'production_order', 'production_order__product', 'inspector'
        ).all().order_by('-inspected_at')[:100]
    else:
        # พนักงานทั่วไป: เห็นเฉพาะประวัติที่ "ตัวเอง" เป็นคนตรวจ
        qc_logs = QCInspectionLog.objects.select_related(
            'production_order', 'production_order__product', 'inspector'
        ).filter(inspector=emp).order_by('-inspected_at')[:100]

    return render(request, 'manufacturing/my_qc_log.html', {
        'qc_logs': qc_logs,
        'is_manager': is_manager, # ส่งตัวแปรนี้ไปที่หน้าเว็บ
    })

# ==========================================
# 🌟 [NEW] สมองกล: กระดานคิวงานแจ้งซ่อม (Service Board สำหรับช่าง) 🌟
# ==========================================
from .models import ServiceTicket

@login_required
def service_board(request):
    # ดึงตั๋วแจ้งซ่อมทั้งหมดเรียงตามวันล่าสุด
    tickets = ServiceTicket.objects.select_related('production_order', 'salesperson', 'production_team').all().order_by('-created_at')

    col1_pending = []     # รอประเมิน / รอเซลส์คอนเฟิร์ม
    col2_inprogress = []  # นัดช่างแล้ว กำลังซ่อม
    col3_completed = []   # ซ่อมเสร็จแล้ว (โชว์แค่ 30 งานล่าสุดพอ)

    completed_count = 0
    for t in tickets:
        if t.status in ['PENDING', 'WAITING_APPROVAL']:
            col1_pending.append(t)
        elif t.status == 'IN_PROGRESS':
            col2_inprogress.append(t)
        elif t.status in ['COMPLETED', 'CANCELLED']:
            if completed_count < 30:
                col3_completed.append(t)
                completed_count += 1

    teams = ProductionTeam.objects.all().order_by('name')

    return render(request, 'manufacturing/service_board.html', {
        'col1_pending': col1_pending,
        'col2_inprogress': col2_inprogress,
        'col3_completed': col3_completed,
        'teams': teams
    })

@login_required
def process_service_ticket(request, ticket_id):
    ticket = get_object_or_404(ServiceTicket, pk=ticket_id)

    if request.method == 'POST':
        action = request.POST.get('action')

        # 1. หัวหน้าช่างประเมินราคา (กรณีหมดประกัน)
        if action == 'estimate_cost':
            # 🌟 [FIXED] ดักจับกรณีช่างลืมกรอกเลข
            cost_str = request.POST.get('estimated_cost', '0')
            if not cost_str: cost_str = '0'
            cost = cost_str.replace(',', '')

            comment = request.POST.get('cost_comment', '')

            ticket.estimated_cost = Decimal(cost)
            ticket.cost_comment = comment
            ticket.status = 'WAITING_APPROVAL' # เด้งกลับไปให้เซลส์คอนเฟิร์ม
            ticket.save()
            messages.warning(request, f"ส่งประเมินราคา ฿{ticket.estimated_cost:,.2f} ให้เซลส์เพื่อคอนเฟิร์มกับลูกค้าแล้ว")

        # 🌟 [NEW] หัวหน้าช่างอนุโลมซ่อมฟรี (ข้ามขั้นตอนเซลส์) 🌟
        elif action == 'free_repair':
            ticket.estimated_cost = Decimal('0.00')
            ticket.cost_comment = "อนุโลมซ่อมฟรี (หัวหน้าช่างประเมินแล้วไม่มีค่าใช้จ่าย)"
            # ไม่เปลี่ยนสถานะเป็น WAITING_APPROVAL เพื่อให้ช่างจ่ายงานต่อได้เลย
            ticket.save()
            messages.success(request, f"🎁 ดำเนินการอนุโลมซ่อมฟรี! คุณสามารถจ่ายงานให้ทีมช่างเข้าพื้นที่ได้เลยครับ")

        # 2. จ่ายงานให้ทีมช่าง
        elif action == 'assign_team':
            team_id = request.POST.get('production_team')
            apt_date = request.POST.get('appointment_date')

            if team_id:
                ticket.production_team_id = team_id
                if apt_date:
                    ticket.appointment_date = parse_date(apt_date)
                ticket.status = 'IN_PROGRESS'
                ticket.save()
                messages.success(request, f"✅ จ่ายงานแจ้งซ่อม {ticket.code} ให้ทีมช่างเข้าพื้นที่เรียบร้อยแล้ว!")

        # 3. ช่างซ่อมเสร็จ ปิดจ๊อบ
        elif action == 'close_ticket':
            after_files = request.FILES.getlist('after_media_files')
            if after_files:
                for f in after_files:
                    from manufacturing.models import ServiceTicketMedia # ดึงโมเดลมาใช้
                    ServiceTicketMedia.objects.create(
                        ticket=ticket,
                        media_type='AFTER',
                        file=f
                    )

                ticket.status = 'COMPLETED'
                ticket.completed_at = timezone.now()
                ticket.save()
                messages.success(request, f"🎉 ปิดจ๊อบงานซ่อม {ticket.code} และบันทึกรูปลง Media Hub สำเร็จ!")
            else:
                messages.error(request, "❌ กรุณาแนบรูปภาพ/วิดีโอหลังซ่อมเสร็จอย่างน้อย 1 ไฟล์เพื่อปิดงาน")

        # 🌟 [NEW] 4. หัวหน้าช่างกดปิดงาน นำออกจากกระดาน (Archive) 🌟
        elif action == 'archive_ticket':
            ticket.status = 'CLOSED'
            ticket.save()
            messages.success(request, f"📦 ปิดงานซ่อม {ticket.code} ถาวร! การ์ดถูกย้ายไปที่ตารางประวัติเรียบร้อยแล้วครับ")

    return redirect('service_board')

# ==========================================
# 🌟 [NEW] ตารางประวัติแจ้งซ่อม และ ตั้งเบิกค่าแรงซ่อม 🌟
# ==========================================
@login_required
def service_history_list(request):
    if not is_sales_authorized(request.user):
        messages.error(request, "❌ บัญชีของคุณไม่มีสิทธิ์เข้าถึง")
        return redirect('dashboard')

    # ดึงตั๋วแจ้งซ่อมทั้งหมดเรียงตามวันล่าสุด (เน้นงานที่ปิดจ๊อบแล้ว)
    tickets = ServiceTicket.objects.select_related('production_order', 'salesperson', 'production_team').all().order_by('-created_at')

    search_q = request.GET.get('q', '')
    status_filter = request.GET.get('status', '')

    # 🌟 [NEW] รับค่าตัวกรองใหม่
    search_team = request.GET.get('team', '')
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')

    if search_q:
        tickets = tickets.filter(
            Q(code__icontains=search_q) |
            Q(production_order__customer_name__icontains=search_q) |
            Q(production_order__code__icontains=search_q)
        )
    if status_filter:
        tickets = tickets.filter(status=status_filter)

    # 🌟 [NEW] กรองตามทีมช่างที่ประกอบงาน (อ้างอิงจากใบสั่งผลิตต้นทาง)
    if search_team:
        tickets = tickets.filter(production_order__production_team_id=search_team)

    # 🌟 [NEW] กรองตามช่วงวันที่แจ้งซ่อม
    if start_date:
        tickets = tickets.filter(created_at__date__gte=start_date)
    if end_date:
        tickets = tickets.filter(created_at__date__lte=end_date)

    paginator = Paginator(tickets, 20)
    page_obj = paginator.get_page(request.GET.get('page'))

    # 🌟 [NEW] ดึงรายชื่อทีมช่างทั้งหมดไปแสดงใน Dropdown
    teams = ProductionTeam.objects.all().order_by('name')

    return render(request, 'manufacturing/service_history_list.html', {
        'page_obj': page_obj,
        'search_q': search_q,
        'status_filter': status_filter,
        'search_team': search_team,
        'start_date': start_date,
        'end_date': end_date,
        'status_choices': ServiceTicket.STATUS_CHOICES,
        'teams': teams,
    })

@login_required
@transaction.atomic
def submit_service_labor_claim(request, ticket_id):
    if request.method == 'POST':
        ticket = get_object_or_404(ServiceTicket, pk=ticket_id)
        current_emp = getattr(request.user, 'employee', None)

        # 🌟 ป้องกันการเบิกซ้ำ 🌟
        if ticket.labor_claims.exists():
            messages.error(request, f"❌ ตั๋วแจ้งซ่อม {ticket.code} ถูกตั้งเบิกค่าแรงไปแล้ว ไม่สามารถตั้งเบิกซ้ำได้")
            return redirect('service_history_list')

        try:
            amount_str = request.POST.get('claim_amount', '0').replace(',', '')
            amount = Decimal(amount_str)
            contractor_name = request.POST.get('technician_name', '')
            slip_file = request.FILES.get('claim_slip')

            if amount <= 0:
                messages.error(request, "❌ ยอดเงินเบิกต้องมากกว่า 0 บาท")
                return redirect('service_history_list')
            if not slip_file:
                messages.error(request, "❌ กรุณาแนบรูปบิล/หลักฐานหน้างาน")
                return redirect('service_history_list')

            # สร้างใบเบิกค่าแรงซ่อม และบันทึกลงตะกร้า (UNCLAIMED)
            from manufacturing.models import ServiceLaborClaim # ดึงโมเดลมาใช้
            ServiceLaborClaim.objects.create(
                service_ticket=ticket,
                requester=current_emp,
                technician_name=contractor_name,
                amount=amount,
                slip_image=slip_file,
                status='UNCLAIMED'
            )

            messages.success(request, f"🛒 หยอดบิลค่าแรงซ่อม {ticket.code} ลงตะกร้าเบิกส่วนตัวสำเร็จ!")
        except Exception as e:
            messages.error(request, f"❌ เกิดข้อผิดพลาด: {str(e)}")

    return redirect('service_history_list')