from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.db.models import F, Q
import openpyxl
import json

# 🌟 Import ข้อมูลจากแอปตัวเอง (เฉพาะเรื่องคลังสินค้า)
from .models import SolarProduct, SolarProductCategory, SolarRawMaterialCategory, SolarStockMovement, SolarInventoryDoc
from django.core.paginator import Paginator
from .forms import SolarProductForm, SolarStockMovementForm, SolarStandardBOMFormSet

# 🌟 [NEW] นำเข้าโมเดลใบสั่งงานมาใช้ตัดสต็อก
from solar_jobs.models import SolarJob, SolarJobBOM
# 🌟 [NEW] นำเข้าโมเดลจากฝั่งจัดซื้อเพื่อนับยอดรอรับของ
from solar_purchasing.models import SolarPurchaseOrder

# ==========================================
# 📦 คลังสินค้าโซล่า (Dashboard & แยกลิสต์ FG/RM)
# ==========================================
@login_required
def solar_inventory_list(request):
    rm_products = SolarProduct.objects.filter(product_type='RM', is_active=True)
    fg_products = SolarProduct.objects.filter(product_type='FG', is_active=True)

    low_stock_count = sum(1 for rm in rm_products if rm.stock_qty <= rm.min_level)
    total_rm_value = sum(rm.total_value for rm in rm_products)
    total_fg_value = sum(fg.total_value for fg in fg_products)

    # 🌟 [FIXED] นับรวมงานสถานะ WAITING_PURCHASE (รอจัดซื้อสั่งของ) เข้ามาเป็นงานค้างด้วย
    requisition_count = SolarJob.objects.filter(
        Q(status='WAITING_STORE') |
        Q(status='WAITING_PURCHASE') |
        Q(status='IN_PROGRESS', job_boms__planned_quantity__gt=F('job_boms__actual_used_quantity'))
    ).distinct().count()

    # 🌟 [NEW] นับใบสั่งซื้อ (PO) ที่อนุมัติแล้ว แต่ยังรับของไม่ครบ
    pending_po_count = SolarPurchaseOrder.objects.filter(status='APPROVED').exclude(receipt_status='COMPLETED').count()

    # 🌟 [NEW] แยกดึงประวัติเอกสารเป็น 2 กอง กองละ 10 รายการล่าสุด (โคลนมาจากระบบบ้านน็อคดาวน์)
    recent_rm_docs = SolarInventoryDoc.objects.filter(movements__product__product_type='RM').distinct().order_by('-created_at')[:10]
    recent_fg_docs = SolarInventoryDoc.objects.filter(movements__product__product_type='FG').distinct().order_by('-created_at')[:10]

    return render(request, 'solar_inventory/dashboard.html', {
        'rm_count': rm_products.count(),
        'fg_count': fg_products.count(),
        'low_stock_count': low_stock_count,
        'total_rm_value': total_rm_value,
        'total_fg_value': total_fg_value,
        'requisition_count': requisition_count,
        'pending_po_count': pending_po_count,
        'recent_rm_docs': recent_rm_docs,  # 🌟 ส่งประวัติ RM ไปหน้าเว็บ
        'recent_fg_docs': recent_fg_docs   # 🌟 ส่งประวัติ FG ไปหน้าเว็บ
    })

# 🌟 2. หน้าแสดงตาราง แพ็กเกจหลัก (FG) เท่านั้น
@login_required
def solar_fg_list(request):
    fg_products = SolarProduct.objects.filter(product_type='FG').order_by('-is_active', '-created_at')
    return render(request, 'solar_inventory/fg_list.html', {'fg_products': fg_products})

# 🌟 3. หน้าแสดงตาราง อุปกรณ์เสริม (RM) เท่านั้น (อัปเกรดระบบค้นหา)
@login_required
def solar_rm_list(request):
    # 🌟 1. รับค่าจากฟอร์มค้นหา
    search_query = request.GET.get('q', '')
    rm_category_id = request.GET.get('rm_category', '')
    stock_status = request.GET.get('stock_status', '')

    # 🌟 2. ดึงข้อมูลพื้นฐานมาทั้งหมดก่อน
    rm_products = SolarProduct.objects.filter(product_type='RM').order_by('-is_active', '-created_at')

    # 🌟 3. ลอจิกคัดกรองข้อมูล (Filter)
    if search_query:
        rm_products = rm_products.filter(
            Q(code__icontains=search_query) |
            Q(name__icontains=search_query)
        )
    if rm_category_id:
        rm_products = rm_products.filter(rm_category_id=rm_category_id)

    if stock_status == 'in_stock':
        rm_products = rm_products.filter(stock_qty__gt=0)
    elif stock_status == 'out_of_stock':
        rm_products = rm_products.filter(stock_qty__lte=0)

    # 🌟 4. ดึงหมวดหมู่ทั้งหมดไปแสดงใน Dropdown
    rm_categories = SolarRawMaterialCategory.objects.all().order_by('name')

    return render(request, 'solar_inventory/rm_list.html', {
        'rm_products': rm_products,
        'search_query': search_query,
        'rm_categories': rm_categories,
        'selected_rm_cat': rm_category_id,
        'stock_status': stock_status
    })

@login_required
def solar_product_create(request):
    default_type = request.GET.get('type', 'FG')
    page_title = 'เพิ่มวัตถุดิบ/อุปกรณ์เสริม' if default_type == 'RM' else 'เพิ่มสินค้า/แพ็กเกจใหม่'

    rm_prices = {str(rm.id): float(rm.cost_price) for rm in SolarProduct.objects.filter(product_type='RM', is_active=True)}
    rm_prices_json = json.dumps(rm_prices)

    # 🌟 [NEW] สวิตช์สลับราง: เลือกไฟล์ Template ตามประเภทที่รับมา
    template_name = 'solar_inventory/product_fg_form.html' if default_type == 'FG' else 'solar_inventory/product_rm_form.html'

    if request.method == 'POST':
        form = SolarProductForm(request.POST, request.FILES)
        if form.is_valid():
            prod = form.save(commit=False)
            prod.product_type = default_type # 🌟 บังคับเซ็ตค่าประเภทสินค้าเพื่อความชัวร์ (เพราะเราซ่อนช่องนี้ไว้)
            prod.save()

            formset = SolarStandardBOMFormSet(request.POST, instance=prod)
            if formset.is_valid() and default_type == 'FG':
                formset.save()

            messages.success(request, f"✅ เพิ่มรายการ '{prod.name}' ลงในคลังสินค้าเรียบร้อยแล้ว")
            return redirect('solar_inventory_list')
        else:
            messages.error(request, "❌ กรุณาตรวจสอบความถูกต้องของข้อมูล")
            formset = SolarStandardBOMFormSet(request.POST)
    else:
        form = SolarProductForm(initial={'product_type': default_type})
        formset = SolarStandardBOMFormSet()

    return render(request, template_name, {
        'form': form,
        'formset': formset,
        'default_type': default_type,
        'title': page_title,
        'rm_prices_json': rm_prices_json
    })

@login_required
def solar_product_edit(request, pk):
    product = get_object_or_404(SolarProduct, pk=pk)
    p_type = product.product_type # เก็บประเภทสินค้าปัจจุบันไว้

    rm_prices = {str(rm.id): float(rm.cost_price) for rm in SolarProduct.objects.filter(product_type='RM', is_active=True)}
    rm_prices_json = json.dumps(rm_prices)

    # 🌟 [NEW] สวิตช์สลับราง: เลือกไฟล์ Template ตามประเภทของสินค้าที่กำลังแก้ไข
    template_name = 'solar_inventory/product_fg_form.html' if p_type == 'FG' else 'solar_inventory/product_rm_form.html'

    if request.method == 'POST':
        form = SolarProductForm(request.POST, request.FILES, instance=product)
        formset = SolarStandardBOMFormSet(request.POST, instance=product)
        if form.is_valid() and (p_type != 'FG' or formset.is_valid()):
            form.save()
            if p_type == 'FG':
                formset.save()
            messages.success(request, f"✅ อัปเดตข้อมูล '{product.name}' เรียบร้อยแล้ว")
            return redirect('solar_inventory_list')
        else:
            messages.error(request, "❌ กรุณาตรวจสอบความถูกต้องของข้อมูล")
    else:
        form = SolarProductForm(instance=product)
        formset = SolarStandardBOMFormSet(instance=product)

    return render(request, template_name, {
        'form': form,
        'formset': formset,
        'default_type': p_type,
        'product': product,
        'title': f'แก้ไข: {product.name}',
        'rm_prices_json': rm_prices_json
    })

# ------------------------------------------
# 🌟 API สำหรับเพิ่มหมวดหมู่แบบ Popup (AJAX)
# ------------------------------------------
@login_required
def add_fg_category_ajax(request):
    if request.method == 'POST':
        category_name = request.POST.get('name')
        if category_name:
            cat, created = SolarProductCategory.objects.get_or_create(name=category_name.strip())
            return JsonResponse({'success': True, 'id': cat.id, 'name': cat.name})
        return JsonResponse({'success': False, 'error': 'กรุณาระบุชื่อหมวดหมู่'})
    return JsonResponse({'success': False, 'error': 'Invalid request'})

@login_required
def add_rm_category_ajax(request):
    if request.method == 'POST':
        category_name = request.POST.get('name')
        if category_name:
            cat, created = SolarRawMaterialCategory.objects.get_or_create(name=category_name.strip())
            return JsonResponse({'success': True, 'id': cat.id, 'name': cat.name})
        return JsonResponse({'success': False, 'error': 'กรุณาระบุชื่อหมวดหมู่'})
    return JsonResponse({'success': False, 'error': 'Invalid request'})

# ==========================================
# 🌟 [UPDATE] อัปเกรดเทมเพลต และ ระบบนำเข้าให้รองรับ Supplier
# ==========================================
@login_required
def solar_inventory_download_template(request):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="inventory_template.xlsx"'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Solar Inventory"

    # 🌟 [UPDATE] เพิ่มคอลัมน์ 'ชื่อร้านค้า (Supplier)'
    headers = ['ประเภท (FG หรือ RM)', 'รหัสสินค้า (เว้นว่างได้)', 'หมวดหมู่ (Category)', 'ชื่อสินค้า/แพ็กเกจ', 'หน่วยนับ', 'ราคาทุน', 'ราคาขาย', 'จำนวนคงเหลือ', 'จุดสั่งซื้อ', 'ชื่อร้านค้า (Supplier)']
    ws.append(headers)
    ws.append(['RM', 'RM-001', 'อุปกรณ์สายไฟ', 'สายไฟ DC 4mm (ตัวอย่าง)', 'เมตร', 15, 25, 100, 20, 'สยามเคเบิล จำกัด'])
    ws.append(['FG', '', 'แพ็กเกจ 5kW', 'แพ็กเกจ 5kW 1 Phase (ตัวอย่าง)', 'ชุด', 100000, 150000, 5, 2, ''])
    wb.save(response)
    return response

@login_required
def solar_inventory_import(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        if not excel_file.name.endswith('.xlsx'):
            messages.error(request, '❌ ไฟล์ไม่ถูกต้อง กรุณาอัปโหลดไฟล์นามสกุล .xlsx เท่านั้น')
            return redirect('solar_inventory_list')

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            sheet = wb.active
            count_created, count_updated = 0, 0

            # นำเข้าตาราง Supplier จากแอป master_data
            from master_data.models import Supplier

            for row in sheet.iter_rows(min_row=2, values_only=True):
                product_type = str(row[0]).strip().upper() if row[0] else 'RM'
                code = str(row[1]).strip() if row[1] else None
                cat_name = str(row[2]).strip() if row[2] else None
                name = str(row[3]).strip() if row[3] else None
                unit = str(row[4]).strip() if row[4] else 'ชิ้น'
                cost_price = float(row[5]) if row[5] else 0.0
                sell_price = float(row[6]) if row[6] else 0.0
                stock_qty = float(row[7]) if row[7] else 0.0
                min_level = float(row[8]) if row[8] else 0.0

                # 🌟 [NEW] อ่านค่าคอลัมน์ชื่อร้านค้า (คอลัมน์ J)
                supplier_name = str(row[9]).strip() if len(row) > 9 and row[9] else None

                if name and name != 'None':
                    if product_type not in ['FG', 'RM']: product_type = 'RM'

                    fg_cat, rm_cat, supplier_obj = None, None, None

                    if cat_name and cat_name != 'None':
                        if product_type == 'FG': fg_cat, _ = SolarProductCategory.objects.get_or_create(name=cat_name)
                        else: rm_cat, _ = SolarRawMaterialCategory.objects.get_or_create(name=cat_name)

                    # 🌟 [NEW] แปลงชื่อร้านค้าให้กลายเป็นฐานข้อมูล Supplier
                    if supplier_name and supplier_name != 'None':
                        supplier_obj, _ = Supplier.objects.get_or_create(name=supplier_name)

                    defaults_data = {
                        'product_type': product_type, 'name': name, 'category': fg_cat, 'rm_category': rm_cat,
                        'unit': unit, 'cost_price': cost_price, 'sell_price': sell_price, 'stock_qty': stock_qty,
                        'min_level': min_level, 'is_active': True,
                        'supplier': supplier_obj # 🌟 แปะข้อมูลร้านค้าลงไป
                    }
                    if code and code != 'None':
                        obj, created = SolarProduct.objects.update_or_create(code=code, defaults=defaults_data)
                        if created: count_created += 1
                        else: count_updated += 1
                    else:
                        SolarProduct.objects.create(**defaults_data)
                        count_created += 1

            messages.success(request, f'✅ นำเข้าข้อมูลสำเร็จ! สร้างใหม่ {count_created} รายการ, อัปเดต {count_updated} รายการ')
        except Exception as e:
            messages.error(request, f'❌ เกิดข้อผิดพลาดในการอ่านไฟล์: {str(e)}')
    return redirect('solar_inventory_list')

# ------------------------------------------
# 📥/📤 ระบบรับเข้า - เบิกออก แมนนวล (แยกหน้าเด็ดขาด)
# ------------------------------------------
@login_required
def solar_stock_in_create(request):
    if request.method == 'POST':
        form = SolarStockMovementForm(request.POST)
        if form.is_valid():
            movement = form.save(commit=False)
            movement.movement_type = 'IN' # 🌟 บังคับเซ็ตเป็นรับเข้า 100% ป้องกันคนส่งค่าแฮกฟอร์ม

            # สร้างหัวเอกสาร GR อัตโนมัติ
            doc = SolarInventoryDoc.objects.create(
                doc_type='GR',
                reference=movement.reference_doc,
                description="บันทึกรับเข้าคลัง (Manual)",
                created_by=request.user
            )

            movement.doc = doc
            movement.save()

            messages.success(request, f"✅ บันทึกรับเข้า และสร้างเอกสารเลขที่ {doc.doc_no} เรียบร้อยแล้ว")
            return redirect('solar_document_list_in') # บันทึกเสร็จเด้งไปหน้ารวมประวัติ GR
        else:
            messages.error(request, "❌ กรุณาตรวจสอบความถูกต้องของข้อมูล")
    else:
        initial_data = {'movement_type': 'IN'}
        product_id = request.GET.get('product_id')
        if product_id:
            initial_data['product'] = product_id
        form = SolarStockMovementForm(initial=initial_data)

    return render(request, 'solar_inventory/stock_in_form.html', {'form': form})

@login_required
def solar_stock_out_create(request):
    if request.method == 'POST':
        form = SolarStockMovementForm(request.POST)
        if form.is_valid():
            movement = form.save(commit=False)
            movement.movement_type = 'OUT' # 🌟 บังคับเซ็ตเป็นเบิกออก 100%

            # สร้างหัวเอกสาร GI อัตโนมัติ
            doc = SolarInventoryDoc.objects.create(
                doc_type='GI',
                reference=movement.reference_doc,
                description="บันทึกเบิกออก (Manual)",
                created_by=request.user
            )

            movement.doc = doc
            movement.save()

            messages.success(request, f"✅ บันทึกเบิกออก และสร้างเอกสารเลขที่ {doc.doc_no} เรียบร้อยแล้ว")
            return redirect('solar_document_list_out') # บันทึกเสร็จเด้งไปหน้ารวมประวัติ GI
        else:
            messages.error(request, "❌ กรุณาตรวจสอบความถูกต้องของข้อมูล")
    else:
        initial_data = {'movement_type': 'OUT'}
        product_id = request.GET.get('product_id')
        if product_id:
            initial_data['product'] = product_id
        form = SolarStockMovementForm(initial=initial_data)

    return render(request, 'solar_inventory/stock_out_form.html', {'form': form})

# ------------------------------------------
# 📊 ระบบ Stock Card (ประวัติความเคลื่อนไหวสินค้าพร้อม Running Balance)
# ------------------------------------------
@login_required
def solar_stock_card(request, pk):
    product = get_object_or_404(SolarProduct, pk=pk)

    # 🌟 [NEW] ดึงประวัติทั้งหมดโดยเรียงจาก 'เก่าสุด' ไป 'ใหม่สุด' เพื่อคำนวณยอดสะสม
    movements_asc = SolarStockMovement.objects.filter(product=product).order_by('created_at', 'id')

    # 🌟 [NEW] ลอจิกคำนวณยอดคงเหลือสะสม (Running Balance)
    running_balance = 0
    calculated_movements = []

    for move in movements_asc:
        if move.movement_type == 'IN':
            running_balance += move.quantity
        elif move.movement_type == 'OUT':
            running_balance -= move.quantity

        # แปะค่ายอดสะสม (running_balance) ไว้ที่ object แต่ละตัว
        move.current_balance = running_balance
        calculated_movements.append(move)

    # 🌟 [NEW] กลับด้านข้อมูล (Reverse) ให้รายการ 'ล่าสุด' ขึ้นมาอยู่บนสุดเหมือนเดิม
    calculated_movements.reverse()

    # นำรายการที่คำนวณเสร็จแล้วไปแบ่งหน้า (Pagination)
    paginator = Paginator(calculated_movements, 20)
    page_obj = paginator.get_page(request.GET.get('page'))

    # 🌟 [NEW] สวิตช์สลับราง: ตรวจสอบประเภทสินค้าเพื่อเลือกไฟล์หน้าจอให้ถูกต้อง
    if product.product_type == 'FG':
        template_name = 'solar_inventory/stock_card_fg.html'
    else:
        template_name = 'solar_inventory/stock_card_rm.html'

    return render(request, template_name, {
        'product': product,
        'page_obj': page_obj
    })

# ==========================================
# 📦 หน้าจอสำหรับสโตร์: จัดการใบขอเบิกวัสดุ (Requisitions)
# ==========================================
@login_required
def store_requisition_list(request):
    # 🌟 [FIXED] ปรับสูตรให้ดึงงานที่รอสโตร์, เบิกเพิ่ม หรือ "รอจัดซื้อ" มาแสดงทั้งหมด (ให้ตรงกับ Dashboard)
    jobs_waiting = SolarJob.objects.filter(
        Q(status='WAITING_STORE') |
        Q(status='WAITING_PURCHASE') |
        Q(status='IN_PROGRESS', job_boms__planned_quantity__gt=F('job_boms__actual_used_quantity'))
    ).prefetch_related('job_boms__product').distinct().order_by('-created_at')

    return render(request, 'solar_inventory/store_requisition_list.html', {
        'jobs_waiting': jobs_waiting
    })

# 🌟 [FIXED] ฟังก์ชันแสดงหน้า Detail พร้อมระบบคำนวณของขาดและปลดล็อกอัตโนมัติ
@login_required
def store_requisition_detail(request, job_id):
    job = get_object_or_404(SolarJob, id=job_id)

    has_shortage = False
    bom_list = []

    # 🌟 ลูปที่ 1: คำนวณหาของขาด (เช็คสต็อก Real-time)
    for bom in job.job_boms.all():
        qty_to_deduct = bom.planned_quantity - bom.actual_used_quantity
        missing_qty = 0

        if qty_to_deduct > 0:
            if qty_to_deduct > bom.product.stock_qty:
                missing_qty = qty_to_deduct - bom.product.stock_qty
                has_shortage = True # เจอของขาด
            bom.qty_to_deduct = qty_to_deduct
            bom.missing_qty = missing_qty
            bom_list.append(bom)

    # 🌟 [NEW] ลอจิกปลดล็อกอัตโนมัติ (Auto-Unlock)
    # ถ้าสถานะเดิมถูกล็อกไว้ว่ารอจัดซื้อ แต่ตอนนี้สต็อกตรวจพบว่า "ของครบแล้ว" (ไม่มี has_shortage)
    if (job.status == 'WAITING_PURCHASE' or job.is_waiting_purchase) and not has_shortage:
        job.is_waiting_purchase = False
        job.status = 'WAITING_STORE'
        job.save()

        # 🌟 ล้างสถานะ PPO ในฝั่งจัดซื้อที่ผูกกับใบนี้ด้วย (ถ้ามี) ให้กลายเป็นอดีตไป
        from solar_purchasing.models import SolarPurchasePreparation
        ppo = SolarPurchasePreparation.objects.filter(solar_jobs=job).exclude(status='CANCELLED').first()
        if ppo:
            ppo.status = 'COMPLETED'
            ppo.save()

        messages.success(request, f"✨ ออโต้เมชั่น: สินค้าครบแล้ว! ระบบทำการปลดล็อกให้จ่ายของได้อัตโนมัติ")

    return render(request, 'solar_inventory/store_requisition_detail.html', {
        'job': job,
        'bom_list': bom_list,
        'has_shortage': has_shortage
    })

# 🌟 [FIXED] ฟังก์ชันสำหรับให้สโตร์กด "ส่งใบขอซื้อ (PR)" ให้จัดซื้อ (ดึงตารางจากแอปจัดซื้อที่ถูกต้อง)
@login_required
def store_trigger_pr(request, job_id):
    if request.method == 'POST':
        job = get_object_or_404(SolarJob, id=job_id)

        # 🌟 [UPDATE] นำเข้าตาราง PPO จากแอป solar_purchasing โดยตรง
        from solar_purchasing.models import SolarPurchasePreparation, SolarPurchasePreparationItem

        shortage_items = []
        for bom in job.job_boms.all():
            qty_to_deduct = bom.planned_quantity - bom.actual_used_quantity
            if qty_to_deduct > 0 and qty_to_deduct > bom.product.stock_qty:
                missing_qty = qty_to_deduct - bom.product.stock_qty
                shortage_items.append({
                    'product': bom.product,
                    'missing_qty': missing_qty
                })

        if not shortage_items:
            messages.warning(request, "⚠️ ไม่พบรายการวัสดุที่ขาด ไม่สามารถส่งเรื่องขอซื้อได้")
            return redirect('store_requisition_list')

        # 🌟 สร้างใบ PPO เปล่าขึ้นมา
        ppo = SolarPurchasePreparation.objects.create(
            created_by=getattr(request.user, 'employee', None),
            status='PENDING'
        )

        # 🌟 [UPDATE] ผูกงาน (JOB) เข้าไปในตะกร้า PPO (เนื่องจากโครงสร้างใหม่เป็น ManyToMany)
        ppo.solar_jobs.add(job)

        # สร้างรายการวัสดุที่ขาด
        for item in shortage_items:
            SolarPurchasePreparationItem.objects.create(
                ppo=ppo,
                product=item['product'],
                quantity_needed=item['missing_qty'],
                ordered_qty=0
            )

        if job.status != 'IN_PROGRESS':
            job.status = 'WAITING_PURCHASE'

        job.is_waiting_purchase = True
        job.save()

        messages.success(request, f"✅ ระบบได้สร้างใบขอซื้อ (PPO) รหัส {ppo.code} พร้อมดึงรายการวัสดุที่ขาดส่งไปให้แผนกจัดซื้อเรียบร้อยแล้ว!")
        return redirect('store_requisition_list')

    return redirect('store_requisition_list')

@login_required
def store_confirm_deduction(request, job_id):
    if request.method == 'POST':
        job = get_object_or_404(SolarJob, id=job_id)

        # 🌟 ลูปที่ 1: เช็คก่อนว่ามีของชิ้นไหนที่สต็อกไม่พอหรือจะทำให้ติดลบหรือไม่
        has_items_to_deduct = False
        for bom in job.job_boms.all():
            qty_to_deduct = bom.planned_quantity - bom.actual_used_quantity
            if qty_to_deduct > 0:
                has_items_to_deduct = True
                if bom.product and bom.product.stock_qty < qty_to_deduct:
                    messages.error(request, f"❌ ไม่สามารถจ่ายของได้! วัตถุดิบ '{bom.product.name}' มีจำนวนไม่เพียงพอ (ต้องการเบิก {qty_to_deduct}, มีอยู่ {bom.product.stock_qty})")
                    return redirect('store_requisition_list')

        # 🌟 [NEW] สร้างหัวเอกสารใบเบิก (GI) 1 ใบคลุมรายการทั้งหมดที่จะจ่าย
        gi_doc = None
        if has_items_to_deduct:
            gi_doc = SolarInventoryDoc.objects.create(
                doc_type='GI',
                reference=f"JOB: {job.code}",
                description=f"เบิกจ่ายวัตถุดิบ/อุปกรณ์ สำหรับงานติดตั้ง {job.code}",
                created_by=request.user
            )

            # 🌟 ลูปที่ 2: ถ้ารอดเงื่อนไขด้านบนมาได้ แสดงว่าของครบ ค่อยดึงรายการเพื่อตัดสต็อกจริง
            for bom in job.job_boms.all():
                qty_to_deduct = bom.planned_quantity - bom.actual_used_quantity

                if qty_to_deduct > 0 and bom.product:
                    # 🌟 สร้างประวัติเพื่อตัดสต็อก พร้อมผูกรหัสใบเบิก (GI) เข้าไปด้วย
                    SolarStockMovement.objects.create(
                        doc=gi_doc,  # <--- พระเอกของเรา ผูกเข้ากับหัวบิล
                        product=bom.product,
                        quantity=qty_to_deduct,
                        movement_type='OUT',
                        reference_doc=f"จ่ายของสำหรับงาน: {job.code}"
                    )

                    # 🌟 อัปเดตยอดเบิกจริงใน BOM ให้รวมกับยอดที่เพิ่งเบิกไป
                    bom.actual_used_quantity += qty_to_deduct
                    bom.save()

        # 🌟 คืนสถานะงานกลับไปเป็น "กำลังติดตั้ง (IN_PROGRESS)"
        job.status = 'IN_PROGRESS'
        job.save()

        if gi_doc:
            messages.success(request, f"✅ ตัดสต็อกและสร้างใบเบิกเลขที่ {gi_doc.doc_no} สำหรับงาน {job.code} เรียบร้อยแล้ว!")
        else:
            messages.warning(request, f"⚠️ ไม่มีรายการเบิกเพิ่มเติมสำหรับงาน {job.code}")

    return redirect('store_requisition_list')

# ==========================================
# 🌟 [NEW] ระบบจัดการเอกสารคลังสินค้า (Stock In / Print A4) - อัปเกรดระบบค้นหาแบบปลอดภัย (Safe Mode)
# ==========================================

# 🌟 [FIXED] เพิ่มฟังก์ชันพิมพ์เอกสารที่หายไปกลับเข้ามา
@login_required
def solar_print_document(request, doc_no):
    doc = get_object_or_404(SolarInventoryDoc, doc_no=doc_no)

    # ลองดึงข้อมูลบริษัทมาใช้พิมพ์บนหัวกระดาษ (ถ้ามี)
    company = None
    try:
        from master_data.models import CompanyInfo
        company = CompanyInfo.objects.first()
    except ImportError:
        pass

    return render(request, 'solar_inventory/doc_print.html', {
        'doc': doc,
        'company': company
    })

@login_required
def solar_document_list_in(request):
    docs = SolarInventoryDoc.objects.filter(doc_type='GR')

    # 🌟 1. ดึงค่าจากหน้าเว็บแบบปลอดภัย
    search_query = request.GET.get('q', '')
    if search_query:
        search_query = search_query.strip()

    product_type = request.GET.get('product_type', '')
    date_start = request.GET.get('start', '')
    date_end = request.GET.get('end', '')

    # 🌟 2. ค้นหาจาก เลขที่เอกสาร, อ้างอิง หรือ หมายเหตุ
    if search_query:
        docs = docs.filter(
            Q(doc_no__icontains=search_query) |
            Q(reference__icontains=search_query) |
            Q(description__icontains=search_query)
        )

    # 🌟 3. กรองตามประเภทสินค้า
    if product_type:
        docs = docs.filter(movements__product__product_type=product_type).distinct()

    # 🌟 4. กรองตามช่วงวันที่ (ใส่ Try-Except กันระบบรวนหากบราวเซอร์ส่งวันที่ผิด Format)
    try:
        if date_start:
            docs = docs.filter(created_at__date__gte=date_start)
        if date_end:
            docs = docs.filter(created_at__date__lte=date_end)
    except Exception:
        pass

    docs = docs.order_by('-created_at')

    return render(request, 'solar_inventory/document_list_in.html', {
        'docs': docs,
        'title': 'ประวัติใบรับสินค้า/วัตถุดิบ (Stock In)',
        'search_query': search_query,
        'product_type': product_type,
        'date_start': date_start,
        'date_end': date_end
    })

@login_required
def solar_document_list_out(request):
    docs = SolarInventoryDoc.objects.filter(doc_type='GI')

    # 🌟 1. ดึงค่าจากหน้าเว็บแบบปลอดภัย
    search_query = request.GET.get('q', '')
    if search_query:
        search_query = search_query.strip()

    product_type = request.GET.get('product_type', '')
    date_start = request.GET.get('start', '')
    date_end = request.GET.get('end', '')

    # 🌟 2. ค้นหาจาก เลขที่เอกสาร, อ้างอิง หรือ หมายเหตุ
    if search_query:
        docs = docs.filter(
            Q(doc_no__icontains=search_query) |
            Q(reference__icontains=search_query) |
            Q(description__icontains=search_query)
        )

    # 🌟 3. กรองตามประเภทสินค้า
    if product_type:
        docs = docs.filter(movements__product__product_type=product_type).distinct()

    # 🌟 4. กรองตามช่วงวันที่ (เกราะป้องกัน Error)
    try:
        if date_start:
            docs = docs.filter(created_at__date__gte=date_start)
        if date_end:
            docs = docs.filter(created_at__date__lte=date_end)
    except Exception:
        pass

    docs = docs.order_by('-created_at')

    return render(request, 'solar_inventory/document_list_out.html', {
        'docs': docs,
        'title': 'ประวัติใบเบิกสินค้า (Stock Out / GI)',
        'search_query': search_query,
        'product_type': product_type,
        'date_start': date_start,
        'date_end': date_end
    })

# ==========================================
# 🚚 ระบบจัดการคิวรับสินค้าเข้าคลัง (Pending PO) สำหรับสโตร์
# ==========================================
# สร้างฟังก์ชันเช็คสิทธิ์ (กรณีที่ยังไม่มีในไฟล์นี้)
def is_inventory_staff(user):
    if user.is_superuser: return True
    if hasattr(user, 'employee') and user.employee:
        dept = getattr(user.employee.department, 'name', '')
        if 'คลังสินค้า' in dept or 'สโตร์' in dept or 'Inventory' in dept or 'Store' in dept:
            return True
    return False

@login_required
def store_pending_po_list(request):
    if not is_inventory_staff(request.user):
        messages.error(request, "❌ บัญชีของคุณไม่มีสิทธิ์เข้าถึงหน้านี้")
        return redirect('dashboard')

    # ดึงเฉพาะใบสั่งซื้อที่อนุมัติแล้ว แต่ยังรับของไม่ครบ
    pending_pos = SolarPurchaseOrder.objects.filter(status='APPROVED').exclude(receipt_status='COMPLETED').order_by('-created_at')

    return render(request, 'solar_inventory/store_pending_po_list.html', {
        'pending_pos': pending_pos
    })

# ==========================================
# 🌟 [NEW] ฟังก์ชันสำหรับพิมพ์ใบจัดของ (Picking List) และ ดูประวัติเบิกย้อนหลัง
# ==========================================
from django.utils import timezone

@login_required
def store_print_picking_list(request, job_id):
    # 1. ดึงข้อมูลงาน (Job)
    job = get_object_or_404(SolarJob, id=job_id)

    # 2. คัดกรองเฉพาะรายการที่ "ต้องจัดของเพิ่ม" เท่านั้น
    bom_list = []
    for bom in job.job_boms.all():
        qty_to_deduct = bom.planned_quantity - bom.actual_used_quantity
        if qty_to_deduct > 0:
            bom.qty_to_deduct = qty_to_deduct
            bom_list.append(bom)

    return render(request, 'solar_inventory/store_print_picking_list.html', {
        'job': job,
        'bom_list': bom_list,
        'print_date': timezone.now()
    })

@login_required
def store_requisition_history(request):
    # 1. ดึงประวัติใบเบิก (GI) ที่อ้างอิงถึง JOB โซล่าเซลล์เรียงจากใหม่ไปเก่า
    gi_docs = SolarInventoryDoc.objects.filter(
        doc_type='GI',
        reference__icontains='JOB'
    ).order_by('-created_at')

    # 2. นำรายการมาแบ่งหน้า (Pagination) หน้าละ 20 รายการเพื่อไม่ให้โหลดช้า
    from django.core.paginator import Paginator
    paginator = Paginator(gi_docs, 20)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'solar_inventory/store_requisition_history.html', {
        'page_obj': page_obj
    })

@login_required
def store_print_gi_doc(request, doc_no):
    # 1. ดึงหัวเอกสาร GI ตามเลขที่กดเข้ามา
    gi_doc = get_object_or_404(SolarInventoryDoc, doc_no=doc_no)

    # 2. ดึงรายการวัสดุที่ถูกตัดสต็อกออกไป "ในบิลใบนี้"
    movements = SolarStockMovement.objects.filter(doc=gi_doc).order_by('id')

    return render(request, 'solar_inventory/store_print_gi_doc.html', {
        'gi_doc': gi_doc,
        'movements': movements,
        'print_date': timezone.now()
    })

# ==========================================
# 🌟 [NEW] ระบบพิมพ์ใบตรวจนับสต็อก (Stock Count Sheet) 🌟
# ==========================================
@login_required
def solar_print_stock_count(request):
    p_type = request.GET.get('type', 'RM') # ค่าเริ่มต้นให้เป็น RM
    search_query = request.GET.get('q', '')
    rm_category_id = request.GET.get('rm_category', '')
    stock_status = request.GET.get('stock_status', '')

    # 1. ดึงข้อมูลพื้นฐาน
    products = SolarProduct.objects.filter(is_active=True, product_type=p_type).order_by('code')

    # 2. กรองข้อมูลตามที่ส่งมาจากหน้าเว็บ (เพื่อให้ตรงกับที่ตาเห็น)
    if search_query:
        products = products.filter(Q(code__icontains=search_query) | Q(name__icontains=search_query))
    if rm_category_id:
        products = products.filter(rm_category_id=rm_category_id)

    if stock_status == 'in_stock':
        products = products.filter(stock_qty__gt=0)
    elif stock_status == 'out_of_stock':
        products = products.filter(stock_qty__lte=0)

    # 3. ดึงชื่อแผนกมาแสดงที่หัวกระดาษ (ถ้ามีการเลือกกรอง)
    category_name = "ทั้งหมด"
    if rm_category_id:
        try: category_name = SolarRawMaterialCategory.objects.get(id=rm_category_id).name
        except: pass

    # 4. ดึงข้อมูลบริษัท (ถ้ามี)
    company = None
    try:
        from master_data.models import CompanyInfo
        company = CompanyInfo.objects.first()
    except ImportError:
        pass

    title = 'คลังแพ็กเกจหลัก (FG)' if p_type == 'FG' else 'คลังอุปกรณ์เสริม (RM)'

    return render(request, 'solar_inventory/stock_count_print.html', {
        'products': products,
        'title': title,
        'p_type': p_type,
        'category_name': category_name,
        'company': company
    })