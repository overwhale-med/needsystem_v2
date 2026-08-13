import json
from decimal import Decimal
from django.db.models import Sum
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import SolarQuotation, SolarQuotationItem, SolarInvoice, SolarProduct, SolarProductCategory, SolarRawMaterialCategory
from .forms import SolarQuotationStep1Form, SolarProductForm
from solar_jobs.models import SolarJob
import openpyxl
from django.http import HttpResponse
from django.utils.dateparse import parse_date
from django.core.paginator import Paginator
from django.db.models import Q
from hr.models import Department
import datetime
import pytz
from django.views.decorators.csrf import csrf_exempt
from django.core.files.base import ContentFile
from master_data.models import CompanyInfo
import base64

# 🌟 [FIXED] เพิ่มบรรทัดนี้เข้ามา เพื่อให้ระบบรู้จักคำสั่ง timezone 🌟
from django.utils import timezone

# ==========================================
# 📊 Dashboards & Lists
# ==========================================
@login_required
def solar_sales_dashboard(request):
    total_sales = SolarInvoice.objects.filter(status='PAID').aggregate(Sum('grand_total'))['grand_total__sum'] or 0
    pending_balance = SolarInvoice.objects.filter(status='UNPAID').aggregate(Sum('balance_amount'))['balance_amount__sum'] or 0
    draft_qt_count = SolarQuotation.objects.filter(status='DRAFT').count()
    unpaid_inv_count = SolarInvoice.objects.filter(status='UNPAID').count()
    recent_invoices = SolarInvoice.objects.all().order_by('-date', '-id')[:5]

    context = {
        'total_sales': total_sales,
        'pending_balance': pending_balance,
        'draft_qt_count': draft_qt_count,
        'unpaid_inv_count': unpaid_inv_count,
        'recent_invoices': recent_invoices,
    }
    return render(request, 'solar_sales/dashboard.html', context)

@login_required
def solar_quotation_list(request):
    # 🌟 [FIXED] เพิ่ม prefetch_related เพื่อโหลดข้อมูล JOB ล่วงหน้า ลดภาระ Database
    quotations = SolarQuotation.objects.prefetch_related('center_jobs').all().order_by('-date', '-id')

    # 🌟 [NEW] ระบบ Auto-Sync: ซิงค์สถานะจาก Center อัตโนมัติ 🌟
    # ถ้ามีงานไหนในกระดาน Center ที่เสร็จสมบูรณ์แล้ว ให้ปรับสถานะใบเสนอราคาเป็น READY ทันที
    for qt in quotations:
        if qt.status == 'PROCESSING':
            for job in qt.center_jobs.all():
                if job.status == 'COMPLETED':
                    # ยิงคำสั่งอัปเดตเข้าฐานข้อมูลทันทีแบบ 100%
                    SolarQuotation.objects.filter(id=qt.id).update(status='READY')
                    qt.status = 'READY'  # อัปเดตภาพบนหน้าจอทันทีเพื่อให้ปุ่มสีม่วงโชว์
                    break

    # 🌟 ระบบจำกัดสิทธิ์และดึงข้อมูลสาขา
    is_manager = False
    is_supervisor = False
    current_emp = getattr(request.user, 'employee', None)

    if request.user.is_superuser:
        is_manager = True
    elif current_emp:
        rank = current_emp.business_rank.lower() if current_emp.business_rank else ""
        if rank in ['manager', 'director'] or 'manager' in getattr(current_emp.position, 'title', '').lower() or 'บัญชี' in getattr(current_emp.department, 'name', ''):
            is_manager = True
        elif rank == 'supervisor':
            is_supervisor = True

    # โหลดรายชื่อสาขา
    if is_manager:
        departments = list(Department.objects.filter(Q(name__icontains='ทีม') | Q(name__icontains='สาขา')).order_by('name'))
    elif is_supervisor and current_emp.department:
        departments = list(Department.objects.filter(id=current_emp.department.id))
    else:
        departments = []

    for d in departments:
        d.name = d.name.replace('แผนก', '').strip()

    # รับค่าการค้นหา
    search_query = request.GET.get('q', '')
    status_filter = request.GET.get('status', '')
    branch_filter = request.GET.get('branch', '')

    # 1. กรองสาขา
    if branch_filter and (is_manager or is_supervisor):
        quotations = quotations.filter(employee__department_id=branch_filter)
    elif not is_manager and not is_supervisor and current_emp:
        # พนักงานทั่วไปเห็นแค่งานตัวเอง
        quotations = quotations.filter(employee=current_emp)

    # 2. กรองสถานะ
    if status_filter:
        if status_filter == 'PENDING_CLOSING':
            quotations = quotations.filter(status='APPROVED', is_deposit_paid=False)
        elif status_filter == 'PENDING_VERIFY':
            quotations = quotations.filter(status='APPROVED', is_deposit_paid=True, is_deposit_verified=False)
        elif status_filter == 'READY':
            quotations = quotations.filter(status='APPROVED', is_deposit_paid=True, is_deposit_verified=True)
        else:
            quotations = quotations.filter(status=status_filter)

    # 3. กรองคำค้นหา
    if search_query:
        quotations = quotations.filter(Q(code__icontains=search_query) | Q(customer__name__icontains=search_query))

    # 4. กรองวันที่
    date_start = request.GET.get('start_date')
    date_end = request.GET.get('end_date')

    if not date_start or date_start == 'None':
        date_start = (timezone.now().date() - datetime.timedelta(days=29)).strftime('%Y-%m-%d')
    if not date_end or date_end == 'None':
        date_end = timezone.now().date().strftime('%Y-%m-%d')

    if not status_filter:
        quotations = quotations.filter(date__gte=date_start, date__lte=date_end)

    # ระบบแบ่งหน้า (Pagination)
    paginator = Paginator(quotations, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'solar_sales/quotation_list.html', {
        'quotations': page_obj,
        'page_obj': page_obj,
        'search_query': search_query,
        'is_manager': is_manager,
        'is_supervisor': is_supervisor,
        'departments': departments,
        'branch_filter': branch_filter,
        'status_filter': status_filter,
        'date_start': date_start,
        'date_end': date_end,
    })

# ==========================================
# 📝 ระบบใบเสนอราคา (Knockdown Style)
# ==========================================
@login_required
def solar_quotation_create(request):
    if request.method == 'POST':
        form = SolarQuotationStep1Form(request.POST)
        if form.is_valid():
            qt = form.save(commit=False)
            if hasattr(request.user, 'employee'):
                qt.employee = request.user.employee
            qt.save()
            messages.success(request, f"✅ สร้างเอกสาร {qt.code} เรียบร้อย! กรุณาเพิ่มรายการสินค้า")
            return redirect('solar_quotation_edit', qt_id=qt.id)
    else:
        form = SolarQuotationStep1Form()

    return render(request, 'solar_sales/quotation_form.html', {'form': form})

def calculate_solar_totals(qt):
    subtotal = sum(item.amount for item in qt.items.all())
    qt.subtotal = subtotal
    total_before_vat = subtotal - qt.discount + qt.survey_fee

    if qt.vat_type == 'EXCLUDE':
        qt.vat_amount = total_before_vat * Decimal('0.07')
        qt.grand_total = total_before_vat + qt.vat_amount
        qt.subtotal = total_before_vat
    elif qt.vat_type == 'INCLUDE':
        # 🌟 ถอด VAT 7% จากยอดรวม 🌟
        qt.grand_total = total_before_vat
        qt.vat_amount = (total_before_vat * Decimal('7')) / Decimal('107')
        qt.subtotal = total_before_vat - qt.vat_amount
    else:
        qt.vat_amount = Decimal('0.00')
        qt.grand_total = total_before_vat
        qt.subtotal = total_before_vat

    qt.save()

@login_required
def solar_quotation_edit(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)

    if request.method == 'POST':
        if qt.status != 'DRAFT':
            messages.error(request, "❌ ไม่สามารถแก้ไขได้ เนื่องจากเอกสารนี้ถูกอนุมัติหรือล็อกไปแล้ว")
            return redirect('solar_quotation_edit', qt_id=qt.id)

        # 1. จัดการเพิ่มสินค้าหลัก (FG)
        if 'add_item' in request.POST:
            qty = int(request.POST.get('quantity', 1))
            price = Decimal(request.POST.get('price', '0').replace(',', '') or 0)
            product_id = request.POST.get('product_id')
            item_name = request.POST.get('item_name')

            prod = SolarProduct.objects.filter(id=product_id).first() if product_id else None
            SolarQuotationItem.objects.create(
                quotation=qt, product=prod, item_name=item_name or (prod.name if prod else "สินค้าพิเศษ"),
                quantity=qty, unit_price=price, amount=qty * price
            )
            calculate_solar_totals(qt)
            return redirect('solar_quotation_edit', qt_id=qt.id)

        # 2. จัดการเพิ่มของแถม/Upsale (RM)
        elif 'add_upsale' in request.POST:
            qty = int(request.POST.get('upsale_qty', 1))
            price = Decimal(request.POST.get('upsale_price', '0').replace(',', '') or 0)
            desc = request.POST.get('upsale_desc')

            SolarQuotationItem.objects.create(
                quotation=qt, item_name=desc, quantity=qty, unit_price=price, amount=qty * price
            )
            calculate_solar_totals(qt)
            return redirect('solar_quotation_edit', qt_id=qt.id)

        # 3. อัปเดตข้อมูลทั่วไป
        elif 'update_info' in request.POST or 'finish_quote' in request.POST:
            qt.discount = Decimal(request.POST.get('discount', '0').replace(',', '') or 0)
            qt.survey_fee = Decimal(request.POST.get('survey_fee', '0').replace(',', '') or 0)
            qt.payment_terms = request.POST.get('payment_terms', '')
            qt.note = request.POST.get('note', '')

            # 🌟 [FIXED] เพิ่มคำสั่งรับค่า VAT จากหน้าเว็บมาบันทึกด้วย 🌟
            vat_val = request.POST.get('vat_type')
            if vat_val:
                qt.vat_type = vat_val

            calculate_solar_totals(qt)
            if 'finish_quote' in request.POST:
                messages.success(request, f"✅ สร้างใบเสนอราคา {qt.code} เสร็จสมบูรณ์แล้ว! (รอผู้อนุมัติ)")
                return redirect('solar_quotation_list')
            return redirect('solar_quotation_edit', qt_id=qt.id)

    # === ส่งข้อมูล JSON ให้ Select2 ===
    products_list = [{'id': p.id, 'name': p.name, 'code': p.code, 'sell_price': float(p.sell_price), 'category_id': p.category_id} for p in SolarProduct.objects.filter(is_active=True, product_type='FG')]
    upsales_list = [{'id': a.id, 'name': a.name, 'code': a.code, 'sell_price': float(a.sell_price), 'category_id': a.rm_category_id} for a in SolarProduct.objects.filter(is_active=True, product_type='RM')]

    return render(request, 'solar_sales/quotation_edit.html', {
        'qt': qt,
        'main_categories': SolarProductCategory.objects.all(),
        'upsale_categories': SolarRawMaterialCategory.objects.all(),
        'products_json': json.dumps(products_list),
        'upsales_json': json.dumps(upsales_list),
        'item_total': sum(i.amount for i in qt.items.all()),
    })

@login_required
def solar_delete_item(request, item_id):
    item = get_object_or_404(SolarQuotationItem, pk=item_id)
    qt = item.quotation
    if qt.status != 'DRAFT':
        messages.error(request, "❌ ไม่สามารถลบรายการได้ เนื่องจากเอกสารนี้ถูกอนุมัติหรือล็อกไปแล้ว")
        return redirect('solar_quotation_edit', qt_id=qt.id)
    item.delete()
    calculate_solar_totals(qt)
    return redirect('solar_quotation_edit', qt_id=qt.id)

@login_required
def solar_quotation_approve(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)
    qt.status = 'APPROVED'
    qt.save()
    messages.success(request, f"✅ อนุมัติใบเสนอราคา {qt.code} เรียบร้อยแล้ว")

    # 🌟 [FIXED] เปลี่ยนให้ระบบเด้งกลับมาหน้าตารางรายการ 🌟
    return redirect('solar_quotation_list')

@login_required
def solar_quotation_send_to_center(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)

    if not qt.is_deposit_verified:
        messages.error(request, "❌ ไม่สามารถส่งงานได้: กรุณารอให้แผนกบัญชีตรวจสอบและยืนยันสลิปมัดจำก่อนครับ")
        return redirect('solar_quotation_list')

    # 🌟 [FIXED] เปลี่ยนสถานะเป็น 'PROCESSING' (กำลังดำเนินการ) และยกเลิกการสร้าง Invoice อัตโนมัติ 🌟
    qt.status = 'PROCESSING'
    qt.save()

    first_item = qt.items.first()
    package = first_item.product if first_item else None

    # สร้างใบสั่งงาน (JOB) ให้ฝั่ง Center ไปดำเนินการต่อ
    job = SolarJob.objects.create(
        customer=qt.customer,
        salesperson=qt.employee,
        package_sold=package,
        quotation_ref=qt,
        status='DRAFT',
        note=f"📌 สร้างอัตโนมัติจากใบเสนอราคาโซล่า: {qt.code}\nรายละเอียดเพิ่มเติม: {qt.note}"
    )

    messages.success(request, f"🚀 ส่งงานเข้า Center สำเร็จ! ระบบสร้างรหัสงาน {job.code} เรียบร้อยแล้ว (รอช่างดำเนินการติดตั้ง)")
    return redirect('solar_quotation_list')

# ==========================================
# 💰 ระบบการเงิน / Invoice
# ==========================================
@login_required
def solar_invoice_list(request):
    # 🌟 [FIXED] เพิ่ม select_related และ prefetch_related ลดการทำงานซ้ำซ้อนของฐานข้อมูล 🌟
    invoices = SolarInvoice.objects.select_related(
        'customer',
        'quotation_ref',
        'quotation_ref__employee',
        'quotation_ref__employee__department'
    ).prefetch_related('quotation_ref__center_jobs').all().order_by('-date', '-id')

    # 🌟 ระบบจำกัดสิทธิ์และดึงข้อมูลสาขา (อ้างอิงจากผู้เปิดบิล)
    is_manager = False
    is_supervisor = False
    current_emp = getattr(request.user, 'employee', None)

    if request.user.is_superuser:
        is_manager = True
    elif current_emp:
        rank = current_emp.business_rank.lower() if current_emp.business_rank else ""
        if rank in ['manager', 'director'] or 'manager' in getattr(current_emp.position, 'title', '').lower() or 'บัญชี' in getattr(current_emp.department, 'name', ''):
            is_manager = True
        elif rank == 'supervisor':
            is_supervisor = True

    # โหลดรายชื่อสาขา
    if is_manager:
        departments = list(Department.objects.filter(Q(name__icontains='ทีม') | Q(name__icontains='สาขา')).order_by('name'))
    elif is_supervisor and current_emp.department:
        departments = list(Department.objects.filter(id=current_emp.department.id))
    else:
        departments = []

    for d in departments:
        d.name = d.name.replace('แผนก', '').strip()

    # รับค่าการค้นหา
    search_query = request.GET.get('q', '')
    status_filter = request.GET.get('status', '')
    branch_filter = request.GET.get('branch', '')

    # 1. กรองสาขา
    if branch_filter and (is_manager or is_supervisor):
        invoices = invoices.filter(quotation_ref__employee__department_id=branch_filter)
    elif not is_manager and not is_supervisor and current_emp:
        # พนักงานทั่วไปเห็นแค่งานตัวเอง
        invoices = invoices.filter(quotation_ref__employee=current_emp)

    # 2. กรองสถานะ
    if status_filter:
        invoices = invoices.filter(status=status_filter)

    # 3. กรองคำค้นหา (ค้นหาได้ทั้งเลขบิล, ชื่อลูกค้า, และเลข QT)
    if search_query:
        invoices = invoices.filter(
            Q(code__icontains=search_query) |
            Q(customer__name__icontains=search_query) |
            Q(quotation_ref__code__icontains=search_query)
        )

    # 4. กรองวันที่ (ค่าเริ่มต้นคือวันที่ปัจจุบันของไทย)
    tz_bkk = pytz.timezone('Asia/Bangkok')
    today_bkk = timezone.now().astimezone(tz_bkk).date()

    date_start = request.GET.get('start_date')
    date_end = request.GET.get('end_date')

    if not date_start or date_start == 'None':
        date_start = today_bkk.strftime('%Y-%m-%d')
    if not date_end or date_end == 'None':
        date_end = today_bkk.strftime('%Y-%m-%d')

    invoices = invoices.filter(date__gte=date_start, date__lte=date_end)

    # ระบบแบ่งหน้า (Pagination)
    paginator = Paginator(invoices, 15)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'solar_sales/invoice_list.html', {
        'page_obj': page_obj,
        'search_query': search_query,
        'is_manager': is_manager,
        'is_supervisor': is_supervisor,
        'departments': departments,
        'branch_filter': branch_filter,
        'status_filter': status_filter,
        'start_date': date_start,
        'end_date': date_end,
    })

@login_required
def solar_quotation_print(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)

    # ดึงข้อมูลบริษัทมาโชว์ที่หัวกระดาษ
    from master_data.models import CompanyInfo
    company = CompanyInfo.objects.first()

    # คำนวณราคาสินค้ารวมก่อนหักส่วนลด
    item_total = sum(item.amount for item in qt.items.all())

    return render(request, 'solar_sales/quotation_print.html', {
        'qt': qt,
        'company': company,
        'item_total': item_total
    })

@login_required
def solar_invoice_print(request, inv_id):
    inv = get_object_or_404(SolarInvoice, pk=inv_id)
    return render(request, 'solar_sales/invoice_print.html', {'inv': inv})

# ==========================================
# 📦 คลังสินค้าโซล่า (Inventory & Excel Import)
# ==========================================
@login_required
def solar_inventory_list(request):
    fg_products = SolarProduct.objects.filter(product_type='FG').order_by('-is_active', '-created_at')
    rm_products = SolarProduct.objects.filter(product_type='RM').order_by('-is_active', '-created_at')
    return render(request, 'solar_sales/inventory_list.html', {
        'fg_products': fg_products,
        'rm_products': rm_products
    })

@login_required
def solar_product_create(request):
    default_type = request.GET.get('type', 'FG')
    page_title = 'เพิ่มวัตถุดิบ/อุปกรณ์เสริม' if default_type == 'RM' else 'เพิ่มสินค้า/แพ็กเกจใหม่'

    if request.method == 'POST':
        form = SolarProductForm(request.POST)
        if form.is_valid():
            prod = form.save()
            messages.success(request, f"✅ เพิ่มรายการ '{prod.name}' ลงในคลังสินค้าเรียบร้อยแล้ว")
            return redirect('solar_inventory_list')
        else:
            messages.error(request, "❌ กรุณาตรวจสอบความถูกต้องของข้อมูล")
    else:
        form = SolarProductForm(initial={'product_type': default_type})

    return render(request, 'solar_sales/product_form.html', {'form': form, 'title': page_title})

@login_required
def solar_product_edit(request, pk):
    product = get_object_or_404(SolarProduct, pk=pk)
    if request.method == 'POST':
        form = SolarProductForm(request.POST, instance=product)
        if form.is_valid():
            form.save()
            messages.success(request, f"✅ อัปเดตข้อมูล '{product.name}' เรียบร้อยแล้ว")
            return redirect('solar_inventory_list')
        else:
            messages.error(request, "❌ กรุณาตรวจสอบความถูกต้องของข้อมูล")
    else:
        form = SolarProductForm(instance=product)
    return render(request, 'solar_sales/product_form.html', {'form': form, 'product': product, 'title': f'แก้ไข: {product.name}'})

@login_required
def solar_inventory_download_template(request):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="inventory_template.xlsx"'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Solar Inventory"
    headers = ['ประเภท (FG หรือ RM)', 'รหัสสินค้า (เว้นว่างได้)', 'หมวดหมู่ (Category)', 'ชื่อสินค้า/แพ็กเกจ', 'หน่วยนับ', 'ราคาทุน', 'ราคาขาย', 'จำนวนคงเหลือ', 'จุดสั่งซื้อ']
    ws.append(headers)
    ws.append(['RM', 'RM-001', 'อุปกรณ์สายไฟ', 'สายไฟ DC 4mm (ตัวอย่าง)', 'เมตร', 15, 25, 100, 20])
    ws.append(['FG', '', 'แพ็กเกจ 5kW', 'แพ็กเกจ 5kW 1 Phase (ตัวอย่าง)', 'ชุด', 100000, 150000, 5, 2])
    wb.save(response)
    return response

@login_required
def solar_inventory_import(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        if not excel_file.name.endswith('.xlsx'):
            messages.error(request, '❌ ไฟล์ไม่ถูกต้อง กรุณาอัปโหลดไฟล์นามสกุล .xlsx เท่านั้น')
            return redirect('solar_inventory_list')

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            sheet = wb.active
            count_created, count_updated = 0, 0

            for row in sheet.iter_rows(min_row=2, values_only=True):
                product_type = str(row[0]).strip().upper() if row[0] else 'RM'
                code = str(row[1]).strip() if row[1] else None
                cat_name = str(row[2]).strip() if row[2] else None
                name = str(row[3]).strip() if row[3] else None
                unit = str(row[4]).strip() if row[4] else 'ชิ้น'
                cost_price = float(row[5]) if row[5] else 0.0
                sell_price = float(row[6]) if row[6] else 0.0
                stock_qty = float(row[7]) if row[7] else 0.0
                min_level = float(row[8]) if row[8] else 0.0

                if name and name != 'None':
                    if product_type not in ['FG', 'RM']: product_type = 'RM'
                    fg_cat, rm_cat = None, None
                    if cat_name and cat_name != 'None':
                        if product_type == 'FG': fg_cat, _ = SolarProductCategory.objects.get_or_create(name=cat_name)
                        else: rm_cat, _ = SolarRawMaterialCategory.objects.get_or_create(name=cat_name)

                    defaults_data = {
                        'product_type': product_type, 'name': name, 'category': fg_cat, 'rm_category': rm_cat,
                        'unit': unit, 'cost_price': cost_price, 'sell_price': sell_price, 'stock_qty': stock_qty,
                        'min_level': min_level, 'is_active': True
                    }
                    if code and code != 'None':
                        obj, created = SolarProduct.objects.update_or_create(code=code, defaults=defaults_data)
                        if created: count_created += 1
                        else: count_updated += 1
                    else:
                        SolarProduct.objects.create(**defaults_data)
                        count_created += 1

            messages.success(request, f'✅ นำเข้าข้อมูลสำเร็จ! สร้างใหม่ {count_created} รายการ, อัปเดต {count_updated} รายการ')
        except Exception as e:
            messages.error(request, f'❌ เกิดข้อผิดพลาดในการอ่านไฟล์: {str(e)}')
    return redirect('solar_inventory_list')

@login_required
def solar_record_deposit(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)
    if request.method == 'POST':
        amount_str = request.POST.get('deposit_amount', '0').replace(',', '')
        try: amount = Decimal(amount_str)
        except: amount = Decimal('0')

        method = request.POST.get('deposit_method', 'TRANSFER')
        date_str = request.POST.get('deposit_date')

        if amount > 0:
            qt.deposit_amount = amount
            qt.deposit_method = method
            if date_str:
                # ลองแปลงวันที่จากรูปแบบที่ได้รับจาก Flatpickr
                try:
                    qt.deposit_date = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
                except ValueError:
                    qt.deposit_date = timezone.now().date()
            else:
                qt.deposit_date = timezone.now().date()

            qt.is_deposit_paid = True

            # 🌟 [NEW] สมองกลสร้างรหัส RVD-SOL ให้โดยอัตโนมัติ (เหมือนระบบบ้านน็อคดาวน์) 🌟
            if not getattr(qt, 'deposit_code', None):
                tz_bkk = pytz.timezone('Asia/Bangkok')
                now_bkk = timezone.now().astimezone(tz_bkk)
                thai_year = (now_bkk.year + 543) % 100
                prefix = f"RVD-SOL-{thai_year:02d}{now_bkk.strftime('%m')}"

                # หาเลขรันล่าสุดของเดือนนี้ในตารางใบเสนอราคาโซล่า
                last_deposit = SolarQuotation.objects.filter(deposit_code__startswith=prefix).order_by('deposit_code').last()
                seq = int(last_deposit.deposit_code.split('-')[-1]) + 1 if last_deposit else 1
                qt.deposit_code = f"{prefix}-{seq:03d}"

            if 'deposit_slip' in request.FILES:
                qt.deposit_slip = request.FILES['deposit_slip']
            qt.save()
            messages.success(request, f"💰 บันทึกรับมัดจำ {amount:,.2f} บาท และสร้างใบเสร็จ {qt.deposit_code} สำหรับใบเสนอราคา {qt.code} เรียบร้อยแล้ว")
        else:
            messages.error(request, "❌ จำนวนเงินมัดจำต้องมากกว่า 0")

    return redirect('solar_quotation_list')

@login_required
def solar_verify_deposit(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)
    qt.is_deposit_verified = True
    qt.save()
    messages.success(request, f"✅ ยืนยันตรวจสอบยอดมัดจำของ {qt.code} เรียบร้อยแล้ว! งานพร้อมส่งเข้า Center")
    return redirect('solar_quotation_list')

# ==========================================
# 📑 ฟังก์ชันเมนูจัดการ (พิมพ์สัญญา / คัดลอก / ยกเลิก)
# ==========================================
@login_required
def solar_quotation_print_contract(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)

    # ดักไว้ไม่ให้พิมพ์สัญญาถ้ายังไม่มัดจำ
    if not qt.is_deposit_paid:
        messages.error(request, "❌ ไม่สามารถพิมพ์สัญญาได้ เนื่องจากยังไม่มีการบันทึกรับเงินมัดจำ")
        return redirect('solar_quotation_list')

    from master_data.models import CompanyInfo
    company = CompanyInfo.objects.first()
    item_total = sum(item.amount for item in qt.items.all())

    # 🌟 หมายเหตุ: เดี๋ยวเราจะต้องสร้างไฟล์หน้าตาใบสัญญา (quotation_contract.html) มารองรับฟังก์ชันนี้ทีหลังนะครับ
    return render(request, 'solar_sales/quotation_contract.html', {
        'qt': qt,
        'company': company,
        'item_total': item_total
    })

@login_required
def solar_quotation_copy(request, qt_id):
    old_qt = get_object_or_404(SolarQuotation, pk=qt_id)

    # 1. สร้างใบเสนอราคาใหม่ โดยลอกข้อมูลเดิมมา (แต่เปลี่ยนสถานะเป็น DRAFT)
    new_qt = SolarQuotation.objects.create(
        customer=old_qt.customer,
        employee=request.user.employee if hasattr(request.user, 'employee') else old_qt.employee,
        subtotal=old_qt.subtotal,
        discount=old_qt.discount,
        survey_fee=old_qt.survey_fee,
        vat_type=old_qt.vat_type,
        vat_amount=old_qt.vat_amount,
        grand_total=old_qt.grand_total,
        payment_terms=old_qt.payment_terms,
        note=old_qt.note,
        status='DRAFT'
    )

    # 2. คัดลอกรายการสินค้าข้างในมาด้วย
    for item in old_qt.items.all():
        SolarQuotationItem.objects.create(
            quotation=new_qt,
            product=item.product,
            item_name=item.item_name,
            quantity=item.quantity,
            unit_price=item.unit_price,
            amount=item.amount
        )

    messages.success(request, f"✅ คัดลอกเอกสารเรียบร้อยแล้ว ได้เอกสารใหม่เลขที่ {new_qt.code}")
    return redirect('solar_quotation_edit', qt_id=new_qt.id)

@login_required
def solar_quotation_cancel(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)
    qt.status = 'CANCELLED'
    qt.save()

    # 🌟 [NEW] ยกเลิก Job ในกระดาน Center ด้วย (ถ้ามี) 🌟
    # ค้นหาทั้งจาก FK และจากข้อความใน Note เผื่อเป็นงานเก่า
    SolarJob.objects.filter(Q(quotation_ref=qt) | Q(note__icontains=qt.code)).update(status='CANCELLED')

    messages.success(request, f"✅ ยกเลิกเอกสาร {qt.code} และงานในระบบ Center (ถ้ามี) เรียบร้อยแล้ว")
    return redirect('solar_quotation_list')

@login_required
def solar_deposit_list(request):
    # ดึงใบเสนอราคา Solar ที่มีการจ่ายมัดจำแล้ว
    quotations = SolarQuotation.objects.filter(is_deposit_paid=True)

    # 🌟 ระบบจำกัดสิทธิ์ (ผู้จัดการ/บัญชี เห็นทั้งหมด, เซลส์เห็นเฉพาะของตัวเอง)
    is_manager = False
    current_emp = getattr(request.user, 'employee', None)

    if request.user.is_superuser:
        is_manager = True
    elif current_emp:
        rank = current_emp.business_rank.lower() if current_emp.business_rank else ""
        if rank in ['manager', 'director'] or 'manager' in getattr(current_emp.position, 'title', '').lower() or 'บัญชี' in getattr(current_emp.department, 'name', ''):
            is_manager = True

    if not is_manager and current_emp:
        quotations = quotations.filter(employee=current_emp)

    # รับค่าการค้นหาจาก URL
    search_query = request.GET.get('q', '')
    status_filter = request.GET.get('status', '')
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    # กรองคำค้นหา (เลขที่เอกสาร หรือ ชื่อลูกค้า)
    if search_query:
        quotations = quotations.filter(Q(code__icontains=search_query) | Q(customer__name__icontains=search_query))

    # กรองสถานะ
    if status_filter == 'VERIFIED':
        quotations = quotations.filter(is_deposit_verified=True)
    elif status_filter == 'PENDING':
        quotations = quotations.filter(is_deposit_verified=False)

    # 🌟 กรองวันที่มัดจำ 🌟
    if start_date_str and end_date_str:
        try:
            start_date = parse_date(start_date_str)
            end_date = parse_date(end_date_str)
            if start_date and end_date:
                quotations = quotations.filter(deposit_date__range=[start_date, end_date])
        except Exception:
            pass # หากวันที่ผิดรูปแบบให้ปล่อยผ่าน

    # เรียงลำดับและแบ่งหน้า
    quotations = quotations.order_by('-deposit_date', '-id')
    paginator = Paginator(quotations, 15)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'solar_sales/deposit_list.html', {
        'page_obj': page_obj,
        'search_query': search_query,
        'status_filter': status_filter,
        'start_date': start_date_str,
        'end_date': end_date_str,
    })

@login_required
def solar_invoice_detail(request, inv_id):
    inv = get_object_or_404(SolarInvoice, pk=inv_id)
    if request.method == 'POST':
        inv.balance_amount = 0
        inv.status = 'PAID'
        inv.save()
        messages.success(request, f"✅ ยืนยันการรับชำระเงินบิล {inv.code} ปิดยอดเรียบร้อยแล้ว!")
        return redirect('solar_invoice_detail', inv_id=inv.id)
    return render(request, 'solar_sales/invoice_detail.html', {'inv': inv})

# ------------------------------------------
# 🌟 ฟังก์ชันสำหรับปุ่ม "เปิดบิลขาย" (Invoice Creation) 🌟
# ------------------------------------------
@login_required
def solar_quotation_create_invoice(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)

    if qt.status != 'READY':
        messages.error(request, "❌ ไม่สามารถเปิดบิลได้ เอกสารต้องอยู่ในสถานะ 'พร้อมเปิดบิล' เท่านั้น")
        return redirect('solar_quotation_list')

    # เช็คว่ามี Invoice อยู่แล้วหรือไม่ กันเหนียวไม่ให้เปิดบิลซ้ำ
    if hasattr(qt, 'solarinvoice'):
        messages.warning(request, "⚠️ ใบเสนอราคานี้มีการเปิดบิลไปแล้ว")
        return redirect('solar_quotation_list')

    # สร้าง Invoice
    SolarInvoice.objects.create(
        quotation_ref=qt,
        customer=qt.customer,
        grand_total=qt.grand_total,
        balance_amount=qt.grand_total - qt.deposit_amount,
        status='UNPAID' if (qt.grand_total - qt.deposit_amount) > 0 else 'PAID'
    )

    # อัปเดตสถานะใบเสนอราคาเป็น เปิดบิลแล้ว
    qt.status = 'CONVERTED'
    qt.save()

    messages.success(request, f"🎉 สร้างใบเสร็จรับเงินสำหรับ {qt.code} สำเร็จ! สามารถเพิ่มการรับชำระเงินได้ทันที")

    # 🌟 [FIXED] เปลี่ยนให้เด้งไปหน้า "รายการใบเสร็จรับเงิน" ทันที 🌟
    return redirect('solar_invoice_list')

# ==========================================
# ✍️ ระบบเซ็นเอกสารออนไลน์ (Online Signature)
# ==========================================
@csrf_exempt
def solar_customer_sign_quotation(request, token):
    # ดึงใบเสนอราคาตามรหัส Token ลับ
    qt = get_object_or_404(SolarQuotation, signature_token=token)

    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            signature_data = data.get('signature_data')

            if signature_data:
                format, imgstr = signature_data.split(';base64,')
                ext = format.split('/')[-1]
                img_data = base64.b64decode(imgstr)
                file_name = f"sign_QT_SOL_{qt.code}.{ext}"

                qt.customer_signature.save(file_name, ContentFile(img_data), save=False)
                qt.signature_date = timezone.now()
                qt.save()
                return JsonResponse({'success': True})
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})

    item_total = sum(item.amount for item in qt.items.all())
    company = CompanyInfo.objects.first()

    return render(request, 'solar_sales/customer_sign.html', {
        'qt': qt,
        'company': company,
        'item_total': item_total
    })

# 🌟 สำหรับเซ็นสัญญามัดจำ (เตรียมไว้ล่วงหน้าเผื่อใช้งาน)
@csrf_exempt
def solar_customer_sign_deposit(request, token):
    qt = get_object_or_404(SolarQuotation, deposit_signature_token=token)

    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            signature_data = data.get('signature_data')

            if signature_data:
                format, imgstr = signature_data.split(';base64,')
                ext = format.split('/')[-1]
                img_data = base64.b64decode(imgstr)
                file_name = f"sign_DEP_SOL_{qt.code}.{ext}"

                qt.customer_deposit_signature.save(file_name, ContentFile(img_data), save=False)
                qt.deposit_signature_date = timezone.now()
                qt.save()
                return JsonResponse({'success': True})
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})

    company = CompanyInfo.objects.first()
    grand_total = qt.grand_total if qt.grand_total else Decimal('0.00')
    deposit_amount = qt.deposit_amount if qt.deposit_amount else Decimal('0.00')
    balance_due = grand_total - deposit_amount

    # ฟังก์ชันแปลงตัวเลขเป็นอักษรไทย (จำลองมาจากที่เขียนไว้ก่อนหน้า)
    def get_thai_baht_text(number):
        if number == 0: return "ศูนย์บาทถ้วน"
        import math
        number = round(float(number), 2)
        baht = math.floor(number)
        satang = int(round((number - baht) * 100))
        def read_num(n):
            if n == 0: return ""
            numbers = ["", "หนึ่ง", "สอง", "สาม", "สี่", "ห้า", "หก", "เจ็ด", "แปด", "เก้า"]
            positions = ["", "สิบ", "ร้อย", "พัน", "หมื่น", "แสน", "ล้าน"]
            s = str(n)
            length = len(s)
            res = ""
            for i, digit in enumerate(s):
                val = int(digit)
                pos = length - i - 1
                if val == 0: continue
                if pos == 0 and val == 1 and length > 1: res += "เอ็ด"
                elif pos == 1 and val == 1: res += "สิบ"
                elif pos == 1 and val == 2: res += "ยี่สิบ"
                else: res += numbers[val] + positions[pos]
            return res
        res = ""
        if baht > 0: res += read_num(baht) + "บาท"
        if satang > 0: res += read_num(satang) + "สตางค์"
        else: res += "ถ้วน"
        return res

    grand_total_text = get_thai_baht_text(grand_total)
    deposit_amount_text = get_thai_baht_text(deposit_amount)
    balance_due_text = get_thai_baht_text(balance_due)

    # ดึงวันที่จัดส่ง
    job = qt.center_jobs.first()
    delivery_date = job.target_date if job else None

    return render(request, 'solar_sales/customer_sign_deposit.html', {
        'qt': qt,
        'company': company,
        'grand_total_text': grand_total_text,
        'deposit_amount_text': deposit_amount_text,
        'balance_due': balance_due,
        'balance_due_text': balance_due_text,
        'delivery_date': delivery_date,
    })

@login_required
def solar_deposit_print(request, qt_id):
    qt = get_object_or_404(SolarQuotation, pk=qt_id)
    company = CompanyInfo.objects.first()
    balance_due = qt.grand_total - qt.deposit_amount
    return render(request, 'solar_sales/deposit_print.html', {'qt': qt, 'company': company, 'balance_due': balance_due})